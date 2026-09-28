from pathlib import Path
import json,csv,re,unicodedata,base64,html,hashlib
from collections import defaultdict,Counter
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'reports'/'collecte-manuelle'; OUT.mkdir(exist_ok=True)
data=json.loads((ROOT/'reports/LIENS-MANUELS-VERIFIES.json').read_text(encoding='utf-8'))
def plain(s):return re.sub(r'[^a-z0-9]','',unicodedata.normalize('NFKD',s).encode('ascii','ignore').decode().lower())
name_map={}
for x in json.loads((ROOT/'reports/an-pap-2023-links.json').read_text(encoding='utf-8')):
 if 'PAP2023_' in x['url']:
  key=x['url'].split('PAP2023_')[1].replace('.pdf','').replace('CS_','')
  name_map[plain(key)]=x['title']
overrides={
'BG_Gestion_finances_publiques_ressources_humaines':"Gestion des finances publiques et des ressources humaines",
'BG_Investissements_avenir':"Investissements d’avenir",
'BG_Egalite_territoires_logement':"Égalité des territoires et logement",
'BG_Politique_territoires':"Politique des territoires",
'CAS_Transition_energetique':"Transition énergétique",
'CAS_Aides_acquisition_vehicules_propres':"Aides à l’acquisition de véhicules propres",
'CAS_Services_nationaux_transport_conventionnes_voyageurs':"Services nationaux de transport conventionnés de voyageurs",
'CCF_Avances_collectivites_territoriales':"Avances aux collectivités territoriales",
'CCF_Avances_divers_services_Etat_organismes_gerant_services_publics':"Avances à divers services de l’État ou organismes gérant des services publics",
'CCF_Prets_Etats_etrangers':"Prêts à des États étrangers",
'CCF_Prets_avances_particuliers_organismes_prives':"Prêts et avances à des particuliers ou à des organismes privés"}
name_map.update({plain(k):v for k,v in overrides.items()})
pdfs=[]
for r in data['pdfs']:
 x=dict(r);x['title']=name_map.get(plain(r['title']),r['title'])
 x.update(type='PDF direct',status='À récupérer',category=r['title'].split(' ')[0])
 pdfs.append(x)
extra=[
(2017,'Annexe 1 du PLR 2017','https://www.budget.gouv.fr/sites/performance_publique/files/farandole/ressources/2017/rap/pdf/PLR_2017_Annexe1.pdf',162),
(2017,'Annexe 2 du PLR 2017','https://www.budget.gouv.fr/sites/performance_publique/files/farandole/ressources/2017/rap/pdf/PLR_2017_Annexe2.pdf',123),
(2020,'Annexe 1 du PLR 2020','https://www.budget.gouv.fr/documentation/file-download/11146',124),
(2020,'Annexe 2 du PLR 2020','https://www.budget.gouv.fr/documentation/file-download/11149',84),
(2021,'Annexe 1 du PLR 2021','https://www.budget.gouv.fr/documentation/file-download/15358',128),
(2021,'Annexe 2 du PLR 2021','https://www.budget.gouv.fr/documentation/file-download/15361',86),
(2021,'Annexe 1 du PLR 2021 nouvelle présentation','https://www.budget.gouv.fr/documentation/file-download/20829',129),
(2021,'Annexe 2 du PLR 2021 nouvelle présentation','https://www.budget.gouv.fr/documentation/file-download/20841',87),
(2022,'Annexe 1 du PLR 2022','https://www.budget.gouv.fr/documentation/file-download/19743',127),
(2022,'Annexe 2 du PLR 2022','https://www.budget.gouv.fr/documentation/file-download/19746',83),
(2022,'Projet de loi de règlement 2022','https://www.budget.gouv.fr/documentation/file-download/19740',106),
(2023,'RAP complet Aide publique au développement','https://www.budget.gouv.fr/documentation/file-download/22617',159)]
for y,t,u,p in extra:pdfs.append(dict(year=y,title=t,url=u,pages=p,stage='rap',type='PDF direct',status='À récupérer',category='Annexe' if 'Annexe' in t else 'Rapport'))
cats=[dict(r,status='Catalogue',type='Catalogue') for r in data['catalogues']]
an=json.loads((ROOT/'reports/an-dossier-sources.json').read_text(encoding='utf-8'))
for y,u in an.items():cats.append(dict(year=int(y),title='Dossier législatif du règlement à l’Assemblée nationale',url=u,stage='parlement',type='Dossier parlementaire',status='Consultation'))
cats.extend([
dict(year=2017,title='Projet de loi de finances 2017 et procédure parlementaire',url='https://www.assemblee-nationale.fr/14/dossiers/loi_finances_2017.asp',stage='parlement',type='Dossier parlementaire',status='Consultation'),
dict(year=2018,title='Projet de loi de finances 2018 et procédure parlementaire',url='https://www.assemblee-nationale.fr/dyn/15/dossiers/loi_finances_2018',stage='parlement',type='Dossier parlementaire',status='Consultation')])
inventory=list(csv.DictReader((ROOT/'reports/INVENTAIRE-CORPUS.csv').open(encoding='utf-8-sig'),delimiter=';'))
existing=[dict(year=r['exercices'],title=r['titre'],url=r['source_url'],stage=r['categorie'],type=r['format'].upper(),status='Déjà collecté',pages=r['pages_pdf'],sha256=r['sha256'],path=r['chemin_volume']) for r in inventory]
manual=pdfs+cats
records=manual+existing
for i,r in enumerate(records):r['id']=i+1
(OUT/'inventaire-liens.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.7)
sec.left_margin=sec.right_margin=Inches(.75)
for n in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 s=doc.styles[n];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing=1.05
for n,size in [('Title',25),('Heading 1',19),('Heading 2',13)]:
 doc.styles[n].font.size=Pt(size)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(10)
h=sec.header.paragraphs[0];h.text='NOS DENIERS  |  Collecte documentaire'
h.style=doc.styles['Normal'];h.runs[0].font.size=Pt(9)
foot=sec.footer.paragraphs[0];foot.text='Liens vérifiés le 8 septembre 2026  •  '
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
foot.runs[0].font.size=Pt(9)

def link(p,label,url):
 a=OxmlElement('w:hyperlink'); a.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
 c=OxmlElement('w:color');c.set(qn('w:val'),'174E83');pr.append(c)
 u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u)
 r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);a.append(r);p._p.append(a)
 return a
def para(t='',style=None):return doc.add_paragraph(t,style)
def heading(t,level=1):doc.add_heading(t,level)
def newpage(t):doc.add_page_break();heading(t)
def checklink(label,url,suffix=''):
 p=para();p.add_run('☐  ');link(p,label,url)
 if suffix:p.add_run(' — '+suffix)
 return p
def table(headers,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
 for c,w,label in zip(t.rows[0].cells,widths,headers):
  c.width=Inches(w);c.text=label
  for r in c.paragraphs[0].runs:r.bold=True
  sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F1F3F5');c._tc.get_or_add_tcPr().append(sh)
 rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 t.style='Table Grid'
 return t
def style_row(row):
 trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
 for c in row.cells:
  for p in c.paragraphs:
   p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
   for r in p.runs:r.font.size=Pt(10.5)

doc.add_paragraph('Collecter les documents budgétaires de 2017 à 2026','Title')
para('Nos Deniers • Guide de récupération manuelle • 8 septembre 2026','Subtitle')
para('Ce document rassemble les liens confirmés pour compléter le corpus du budget de l’État, toutes missions confondues. La priorité est de récupérer les RAP et PAP de 2017 à 2022, puis les compléments manquants de 2023 à 2026.')
para(f'Le guide comprend {len(pdfs)} liens directs vers des PDF confirmés, ainsi que les catalogues et dossiers par exercice. L’index HTML joint rassemble aussi les références des 767 fichiers déjà collectés. Les pages de catalogue donnent accès à plusieurs documents ; elles ne constituent pas des fichiers ZIP préfabriqués.')
para('La liste des liens directs reste partielle, en particulier pour 2017 et 2018. Les références non confirmées ne sont pas présentées comme des téléchargements valides. Un lien visible dans l’outil de recherche ne garantit pas que le téléchargement automatique fonctionne depuis notre serveur.')
heading('Comment collecter',2)
for t in [
'Ouvrir les liens bleus. Dans Word, utiliser Ctrl + clic si un clic seul ne les ouvre pas. Franchir normalement le contrôle anti-robot si le site le demande.',
'Prendre le PDF complet de chaque mission lorsqu’il existe. Les extraits par programme sont utiles seulement pour compléter un rapport indisponible. Inclure le budget général, les budgets annexes et tous les comptes spéciaux.',
'Dans les catalogues, utiliser la sélection sur toutes les pages et le téléchargement groupé lorsqu’ils sont proposés. Le site affiche une limite de 1 Go : fractionner les lots par année et catégorie.',
'Conserver les noms officiels. Ranger les fichiers par exercice puis PAP, RAP, annexes et compléments ; transmettre les dossiers ou archives obtenus pour l’import.',
'Consulter l’index HTML avant de reprendre un téléchargement récent. L’import final vérifiera les formats et les empreintes, puis isolera les doublons et les versions corrigées.']:
 para(t,'List Number')
para('Exercice et publication sont deux dates différentes : le RAP 2022 retrace 2022, même s’il a été publié en 2023. Un projet de loi de règlement n’est pas automatiquement une loi adoptée.')

for y in [2017,2018]:
 newpage(f'Documents de {y}')
 rows=[r for r in pdfs if r['year']==y and r['category']!='Annexe']
 counts=Counter(r['stage'] for r in rows)
 para(f"{counts['rap']} RAP et {counts['pap']} PAP complets confirmés. Les nombres entre parenthèses sont les pages des PDF. « À retrouver » signifie qu’aucun lien direct n’a été confirmé dans ce relevé, et non que le document n’existe pas.")
 grouped=defaultdict(dict)
 for r in rows:grouped[(r['category'],r['title'])][r['stage']]=r
 tab=table(['Mission ou compte','RAP','PAP'],[4.5,1.25,1.25])
 for (cat,title),pair in sorted(grouped.items(),key=lambda x:(x[0][0],plain(x[0][1]))):
  row=tab.add_row();row.cells[0].text=f'{title} ({cat})'
  for j,s in enumerate(['rap','pap'],1):
   p=row.cells[j].paragraphs[0]
   if s in pair:p.add_run('☐ ');link(p,f"PDF ({pair[s]['pages']} p.)",pair[s]['url'])
   else:p.add_run('À retrouver')
  style_row(row)
 heading('Annexes et accès complémentaires',2)
 for r in pdfs:
  if r['year']==y and r['category']=='Annexe':checklink(r['title'],r['url'],str(r['pages'])+' pages')
 for r in cats:
  if r['year']==y:checklink(r['title'],r['url'])
 checklink('Archives officielles des documents budgétaires 2006 à 2018','https://www.budget.gouv.fr/documents-budgetaires/','navigation dans le site à effectuer manuellement')
 para('À compléter depuis les archives : rapports non confirmés ci-dessus, jaunes budgétaires, documents de politique transversale, voies et moyens, synthèses chiffrées, annexes du règlement et autres comptes spéciaux. Ce relevé n’atteste pas l’exhaustivité de la nomenclature de cet exercice.')

for y in range(2019,2023):
 newpage(f'Catalogues et compléments de {y}')
 para('Les liens ci-dessous ouvrent des ensembles de documents officiels. Télécharger les missions complètes et les annexes de chaque ensemble, sur toutes les pages. Les données tabulaires déjà collectées restent consultables dans l’index joint.')
 yc=[r for r in cats if r['year']==y]
 def rank(r):return (0 if r['stage']=='rap' else 1 if r['stage']=='pap' else 2, r['url'].count('/'),r['title'])
 for r in sorted(yc,key=rank):
  # Mission pages are in the recent section only.
  label=('RAP — ' if r['stage']=='rap' else 'PLF et PAP — ' if r['stage']=='pap' else '')+r['title']
  checklink(label,r['url'])
 annual=[r for r in pdfs if r['year']==y]
 if annual:
  heading('PDF directs confirmés',2)
  for r in annual:checklink(r['title'],r['url'],f"{r['pages']} pages")
 para('Dans le catalogue RAP, prendre aussi la synthèse chiffrée en XLS ou XLSX et les annexes 1 et 2. Dans le catalogue PLF, prendre les jaunes, les documents de politique transversale, les voies et moyens et leurs tableaux annexes. Une même donnée présente dans plusieurs éditions sera conservée avec sa provenance.')

newpage('Compléter les exercices de 2023 à 2026')
para('Les lots récents déjà validés contiennent 47 RAP de mission ou compte pour 2023, 47 pour 2024 et 43 pour 2025, ainsi que 42 PAP pour 2023, 48 pour 2024, 48 pour 2025 et 47 pour 2026. Ces décomptes ne prouvent pas que toutes les missions et annexes sont réunies.')
heading('Rapports et annexes encore à récupérer',2)
for r in cats:
 if r['year']>=2023 and ('/budget-general/' in r['url']):checklink(str(r['year'])+' — '+r['title'],r['url'],'prendre le rapport complet')
for r in pdfs:
 if r['year']>=2023:checklink(str(r['year'])+' — '+r['title'],r['url'],str(r['pages'])+' pages')
for r in cats:
 if r['year']>=2023 and (r['url'].endswith('projet-loi-relatif-aux-resultats') or r['url'].endswith('plrg-2024')):checklink(str(r['year'])+' — Catalogue de tous les RAP et annexes',r['url'])
para('À vérifier dans le catalogue 2025 : Développement agricole et rural, Plan de relance et annexe 2 du PLRG. Leurs liens directs n’ont pas été confirmés ici. Vérifier aussi les PAP récents face aux listes officielles, notamment les comptes de concours financiers de 2023.')
para('Le PDF fourni FR_2025_PLRG_TA_PGM_205.pdf est déjà importé. Il couvre le programme 205 et ne remplace pas le rapport complet de la mission Écologie. Il n’existe pas encore de RAP annuel définitif pour l’exercice 2026 en cours.')
heading('Ce qui doit aussi figurer dans la collecte',2)
para('Inclure les rapports de performance et leurs justificatifs au premier euro, les rapports des opérateurs, les dépenses fiscales, les données de fonds de concours et d’attributions de produits, les tableaux de reports et de mouvements publiés, ainsi que les documents budgétaires rectificatifs et de fin de gestion. Les lois et actes du Journal officiel seront également exploités par API.')

newpage('Téléchargements automatiques et connexions')
para('Les PDF manuels ne sont pas les seuls fichiers nécessaires au service. Les APIs fournissent aussi des données qu’il faut importer et conserver localement pour calculer des séries reproductibles. Une consultation exclusivement à distance ne suffit pas pour l’analyse et l’audit documentaire.')
t=table(['Source','Traitement prévu et situation'],[2,5])
for a,b in [
('data.economie.gouv.fr et data.gouv.fr','Exports et pièces jointes récupérés automatiquement ; première collecte effectuée. Conserver des instantanés datés, puis mettre à jour les séries. Les liens vers un PDF hébergé sur Budget peuvent encore demander une collecte manuelle.'),
('Budget et documents parlementaires','Archiver les PDF et tableaux pour l’extraction, l’indexation et les citations. Certaines archives restent incomplètes malgré les fichiers déjà validés.'),
('Légifrance via PISTE','Accès API testé. Récupérer les lois, décrets et arrêtés utiles aux annulations, ouvertures, transferts, virements et reports. Le rapprochement budgétaire reste à développer.'),
('Insee','Série annuelle d’IPC 2017 à 2025 déjà archivée. Actualisation automatique ; le traitement de 2026 doit distinguer les données disponibles et une moyenne annuelle définitive.'),
('Bases vectorisées existantes','Connexion aux corpus déjà présents ; ne pas les télécharger à nouveau. L’intégration des recherches et la traçabilité des citations restent à réaliser.'),
('Tricoteuses','Données parlementaires publiques accessibles par une API distincte. La connexion OAuth du MCP avait retourné invalid_client ; cette connexion ne doit pas être annoncée comme opérationnelle avant correction et nouveau contrôle.'),
('Cour des comptes et Parlement','Rapports de contrôle et notes d’exécution utiles à l’explication des écarts. À inventorier et importer automatiquement selon l’accès disponible ; ils ne sont pas tous inclus dans le lot actuel.'),
('Réserves et gels','Compléter avec les informations effectivement publiées dans les RAP, les documents parlementaires et les réponses ministérielles. Aucun accès à un flux public exhaustif en temps réel, ventilé par action, n’est établi.'),
('Chorus','Aucun accès interne établi. Les données publiques ne permettent pas de promettre chaque transaction ou une ventilation systématique par sous-action.')]:
 row=t.add_row();row.cells[0].text=a;row.cells[1].text=b;style_row(row)
para('Une valeur absente sera indiquée comme indisponible, jamais assimilée à zéro. Les gels ne sont pas des annulations ; les prévisions de fonds de concours ne sont pas les rattachements réellement constatés.')

docx=OUT/'Nos Deniers - Liens de collecte 2017 à 2026.docx'
for node in (doc._element,doc.styles.element):
 for border in node.xpath('.//w:pBdr'):
  border.getparent().remove(border)
doc.paragraphs[-1].paragraph_format.space_before=Pt(8)
doc.save(docx)

logo=base64.b64encode((ROOT/'assets/nos-deniers/nos-deniers-horizontal.svg').read_bytes()).decode()
payload=json.dumps(records,ensure_ascii=False).replace('<','\\u003c')
page='''<!doctype html><html lang="fr"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Nos Deniers — Tous les liens de collecte</title><style>
*{box-sizing:border-box}body{font:16px system-ui,Arial;color:#173a55;background:#f6f8fb;margin:0}main{max-width:1250px;margin:auto;padding:28px}header img{width:390px;max-width:100%}h1{font-size:27px;color:#08183f}p{line-height:1.55;max-width:1000px}
.filters{display:flex;gap:12px;flex-wrap:wrap;background:white;padding:16px;position:sticky;top:0;border:1px solid #d9e2ed;z-index:1}label{display:grid;gap:5px;font-weight:600;font-size:14px}
input,select,button{font:inherit;padding:10px;border:1px solid #aab9c7;border-radius:4px}input{width:320px;max-width:100%}button{background:#173a55;color:white;cursor:pointer}
table{width:100%;table-layout:fixed;border-collapse:collapse;background:white;font-size:14px}th:first-child{width:70px}th:nth-child(2){width:115px}th:nth-child(3){width:95px}th:nth-child(5){width:140px}td{overflow-wrap:anywhere}th{text-align:left;background:#e7eef5;padding:13px}td{padding:12px;border-bottom:1px solid #e2e8ef;vertical-align:top}
a{color:#174e83;text-decoration:underline;overflow-wrap:anywhere}.status{white-space:nowrap}small{display:block;color:#4c6475;margin-top:5px;overflow-wrap:anywhere}.count{font-weight:600}
details{margin:22px 0}summary{cursor:pointer;font-weight:650}footer{margin-top:25px;font-size:13px}
@media(max-width:700px){main{padding:15px}.filters{position:static}th:nth-child(3),td:nth-child(3){display:none}table{table-layout:fixed}th,td{padding:8px;word-break:break-word}th:first-child,td:first-child{width:40px}th:nth-child(2),td:nth-child(2){width:52px}th:nth-child(5),td:nth-child(5){width:86px}.status{white-space:normal}}
</style><main><header><img src="data:image/svg+xml;base64,LOGO" alt="Nos Deniers — Du budget voté à l’euro dépensé"></header>
<h1>Tous les liens de collecte</h1>
<p>Le répertoire rassemble les liens directs confirmés, les catalogues officiels et les références des <strong>767 fichiers déjà collectés</strong>. Sélectionnez « À récupérer » pour les PDF manquants confirmés, ou « Catalogue » pour les lots disponibles par exercice. Les cases servent de repère pendant cette ouverture ; elles ne sont pas sauvegardées.</p>
<details><summary>Mode d’emploi et limites</summary><p>Les rapports complets d’une mission évitent de télécharger chaque extrait de programme. Incluez aussi les budgets annexes et comptes spéciaux. Dans les catalogues Budget, la sélection peut porter sur toutes les pages, avec une limite affichée de 1 Go par lot. Les archives 2017 et 2018 restent partiellement répertoriées ; aucune adresse supposée n’a été ajoutée comme PDF confirmé.</p>
<p>Les API feront aussi l’objet d’imports automatiques et d’instantanés locaux. Les bases existantes seront connectées. Une ligne « Déjà collecté » décrit le fichier source validé, sans affirmer que ses chiffres sont déjà normalisés ou indexés. L’accès local aux PDF Budget peut être filtré même quand la consultation par le moteur de recherche fonctionne.</p></details>
<div class="filters"><label>Rechercher<input id="q" placeholder="Écologie, annexe, pensions…"></label><label>Exercice<select id="year"><option value="">Tous</option>YEARS</select></label>
<label>État<select id="status"><option value="">Tous</option><option>À récupérer</option><option>Catalogue</option><option>Consultation</option><option>Déjà collecté</option></select></label><button id="reset">Réinitialiser</button></div>
<p id="count" class="count" aria-live="polite"></p><table><thead><tr><th>Repère</th><th>Exercice</th><th>Type</th><th>Document ou ensemble</th><th>État</th></tr></thead><tbody id="rows"></tbody></table>
<footer>Relevé du 8 septembre 2026. Les liens vers des sources publiques ouvrent un nouvel onglet. Aucun téléchargement automatique n’est déclenché par ce fichier.</footer></main>
<script>const DATA=PAYLOAD;const $=id=>document.getElementById(id), norm=x=>String(x).normalize('NFD').replace(/[\\u0300-\\u036f]/g,'').toLowerCase();
function draw(){const q=norm($('q').value),y=$('year').value,s=$('status').value;
const a=DATA.filter(r=>(!q||norm(r.title+' '+r.stage+' '+r.type+' '+r.url).includes(q))&&(!y||String(r.year).includes(y))&&(!s||r.status===s));$('rows').replaceChildren();
for(const r of a){const tr=document.createElement('tr');let td=document.createElement('td');const box=document.createElement('input');box.type='checkbox';box.style.width='18px';box.setAttribute('aria-label','Repérer '+r.title);td.append(box);tr.append(td);
for(const val of [r.year||'Archives',r.type]){td=document.createElement('td');td.textContent=val;tr.append(td)}
td=document.createElement('td');if(/^https?:\\/\\//.test(r.url)){const a=document.createElement('a');a.href=r.url;a.target='_blank';a.rel='noopener noreferrer';a.textContent=(r.stage==='pap'?'PAP — ':r.stage==='rap'?'RAP — ':'')+r.title;td.append(a)}else td.textContent=r.title+' — fichier fourni';
const small=document.createElement('small');small.textContent=(r.pages?r.pages+' pages · ':'')+(r.path||r.url);td.append(small);tr.append(td);td=document.createElement('td');td.textContent=r.status;td.className='status';tr.append(td);$('rows').append(tr)}
$('count').textContent=a.length+' références affichées sur '+DATA.length+' — un catalogue peut donner accès à plusieurs fichiers.'}
for(const id of ['q','year','status'])$(id).addEventListener('input',draw);$('reset').onclick=()=>{for(const id of ['q','year','status'])$(id).value='';draw()};draw();
</script></html>'''
page=page.replace('LOGO',logo).replace('PAYLOAD',payload).replace('YEARS',''.join(f'<option>{y}</option>' for y in range(2017,2027)))
(OUT/'Nos Deniers - Tous les liens.html').write_text(page,encoding='utf-8')
print(json.dumps({'docx':str(docx),'manual_pdf_links':len(pdfs),'catalogues_and_dossiers':len(cats),'all_records':len(records),'old_counts':dict(Counter(str(r['year'])+' '+r['stage'] for r in pdfs if r['year']<2019))},ensure_ascii=False))
