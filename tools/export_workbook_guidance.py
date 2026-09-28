"""Read the explicitly requested guidance sheets without changing the originals."""
from pathlib import Path
import hashlib, json, re
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'reports/reserves-et-consignes-20260909'
SOURCE = Path('C:/Users/Jean-Christophe/Desktop/Nos Deniers')
SHEETS = {
    'Mission écologie 2023 à 2025 .xlsx': ['Explications tableau évolution '],
    'Evolution crédits mission écologie 2017 à 2026 .xlsx': ['Consignes Marie'],
}
records = []
for name, sheets in SHEETS.items():
    path = SOURCE / name
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    wb = load_workbook(path, data_only=False)
    for sheet in sheets:
        ws = wb[sheet]
        cells = []
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None and cell.comment is None:
                    continue
                cells.append(dict(cell=cell.coordinate, value=cell.value,
                    readable_text=re.sub(r'\s+', ' ', str(cell.value or '')).strip(),
                    formula=cell.data_type == 'f',
                    hyperlink=cell.hyperlink.target if cell.hyperlink else None,
                    comment=cell.comment.text if cell.comment else None))
        records.append(dict(workbook=name, path=str(path), sha256=before, sheet=sheet, cells=cells,
            source_kind='user_workbook_guidance', numeric_status='not_validated_budget_fact',
            purpose='Référence métier pour le chantier ; citations et calculs à vérifier dans les publications.',
            previous_export_missing=sheet == 'Consignes Marie'))
    wb.close()
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before
OUT.mkdir(parents=True, exist_ok=True)
(OUT/'explications-excel.json').write_text(json.dumps(records, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
(OUT/'explications-excel.txt').write_text('\n\n'.join(
    f"Classeur : {r['workbook']}\nOnglet : {r['sheet']}\nSHA-256 : {r['sha256']}\n\n"+
    '\n\n'.join(f"Cellule {c['cell']}\n{c['readable_text']}" for c in r['cells']) for r in records), encoding='utf-8')
print(json.dumps([dict(workbook=r['workbook'], sheet=r['sheet'], cells=len(r['cells']), sha256=r['sha256']) for r in records], ensure_ascii=False))
