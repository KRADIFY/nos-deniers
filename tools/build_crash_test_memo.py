from pathlib import Path
import json,datetime
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(__file__).resolve().parents[1];out=root/'reports/crash-test-site-20260924';r=json.loads((out/'report.json').read_text('utf-8'))
cases={x['id']:x for x in r['cases']};vals=[x['result'] for x in r['cases'] if x['status']=='pass']
sumkey=lambda k:sum(x.get(k,0) for x in vals)
num=lambda x:f'{x:,}'.replace(',',' ')
reg=cases['numeric:all-registries']['result'];display=cases.get('browser:display-amounts',{}).get('result',{})
d=Document();sec=d.sections[0];sec.page_height=Cm(29.7);sec.page_width=Cm(21);sec.top_margin=sec.bottom_margin=Cm(1.9);sec.left_margin=sec.right_margin=Cm(2)
for name in ('Normal','Title','Heading 1','Heading 2'):
 st=d.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.space_after=Pt(7)
d.styles['Title'].font.size=Pt(24);d.styles['Heading 1'].font.size=Pt(15)
d.core_properties.title='Nos Deniers Contrôle de solidité et des chiffres';d.core_properties.author='Nos Deniers'
d.add_paragraph('Nos Deniers',style='Title');d.add_paragraph('Contrôle de solidité et des chiffres',style='Subtitle')
d.paragraphs[-1].runs[0].font.color.rgb=RGBColor(0,0,0)
d.add_paragraph('24 septembre 2026 • Version locale testée • Compte rendu pour Jean-Christophe')
fails=r['counts']['failed'];total=r['counts']['total']
d.add_heading('Résultat de la campagne',1)
d.add_paragraph(f"{num(total)} scénarios enregistrés : {num(r['counts']['passed'])} réussis et {fails} en échec. Les contrôles comparent les calculs du site aux données structurées et aux justificatifs enregistrés, puis vérifient les affichages et les parcours dans Chrome. Le site public n’a pas été soumis à cette charge.")
d.add_paragraph('Trois attentes de tests devenues obsolètes ont été corrigées, puis les tests concernés ont été rejoués. Le premier rapport est conservé. La reprise du lanceur a également été corrigée et vérifiée.')
rows=[('Faits présents dans la base',num(sumkey('canonical_rows'))),('Cellules budgétaires comparées',num(sumkey('display_cells'))),('Détails d’actions et sous-actions contrôlés',num(sumkey('derived_action_subaction_cells'))),('Conversions d’inflation contrôlées',num(sumkey('inflation_calculations'))),('Montants de tableau vérifiés à l’écran',num(display.get('visibleAmounts',0))),('Valeurs de graphique vérifiées',num(display.get('chartAmounts',0))),('Groupes d’actions validés dans les registres',num(reg['accepted_action_groups'])),('Observations de réserves contrôlées',num(reg['reserve_observations_checked'])),('Registres historiques de mouvements distincts',num(reg['historical_unique_programme_years']))]
t=d.add_table(rows=1, cols=2);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False;t.columns[0].width=Cm(12);t.columns[1].width=Cm(5)
t.rows[0].cells[0].text='Contrôle';t.rows[0].cells[1].text='Nombre'
for label,value in rows:
 cells=t.add_row().cells;cells[0].text=label;cells[1].text=value
for i,row in enumerate(t.rows):
 for j,c in enumerate(row.cells):
  c.width=Cm(12 if j==0 else 5);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  pr=c._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EEF5' if i==0 else ('F5F7F9' if i%2==0 else 'FFFFFF'));pr.append(shade)
  borders=OxmlElement('w:tcBorders')
  for edge in ('top','left','bottom','right'):
   e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
  pr.append(borders);mar=OxmlElement('w:tcMar')
  for edge in ('top','left','bottom','right'):
   e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');mar.append(e)
  pr.append(mar)
  for p in c.paragraphs:
   p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
   if j==1:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
   for run in p.runs:run.font.size=Pt(10);run.bold=i==0
# Repeat headings, never split an individual row.
for row in t.rows:
 e=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(e)
e=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(e)
d.add_paragraph('Les compteurs représentent des contrôles, pas tous des données distinctes : un montant est notamment testé dans plusieurs unités ou années de référence. La matrice numérique couvre 2017 à 2026, AE et CP, et les quatre catégories budgétaires.')
d.add_heading('Fonctionnement du site',1)
d.add_paragraph('Vérification des filtres, exclusions et réactivations, clics rapprochés, réponses arrivant dans le désordre, données indisponibles, taux et variations, exports CSV et Excel, sources consultables et présentation sur mobile. Des erreurs volontairement introduites dans une copie en mémoire ont permis de vérifier le détecteur.')
loads=[x['result'] for x in r['cases'] if x['kind']=='bounded_load' and x['status']=='pass']
if loads:
 peak=loads[-1];d.add_paragraph(f"Charge locale limitée à {peak['concurrency']} requêtes simultanées. Au dernier palier : {peak['requests']} requêtes, médiane {peak['p50_seconds']:.2f} s, 95e percentile {peak['p95_seconds']:.2f} s, maximum {peak['maximum_seconds']:.2f} s. Les réponses restent identiques à celles obtenues une par une.")
d.add_page_break();d.add_heading('Ce que ce contrôle garantit',1)
d.add_paragraph('Dans les scénarios exécutés, les montants servis et affichés correspondent aux données enregistrées ; les totaux, conversions et taux suivent les formules contrôlées. Les montants indisponibles ne sont pas transformés en zéro. Les tests n’ont modifié ni la base ni la présentation du site.')
d.add_heading('Ce qui reste distinct d’une erreur du site',1)
pending=reg['action_groups_requiring_review'];issues=cases['numeric:source-bindings']['result']['known_source_disagreements']
d.add_paragraph(f"{len(pending)} groupes d’actions restent soumis à examen documentaire. Le contrôle vérifie leur blocage : il ne valide pas leurs ventilations litigieuses. {sumkey('blocked_details_verified')} cellules de détail ont ainsi été contrôlées comme non publiées. Les totaux de programme peuvent rester disponibles.")
d.add_paragraph(f"Les métadonnées conservent également {len(issues)} signalements de source. Ce nombre ne désigne ni {len(issues)} nouveaux bugs ni nécessairement {len(issues)} montants faux. Ces différences et limites sont conservées dans le rapport détaillé ; le test ne les efface pas.")
d.add_paragraph('Les données non isolées de MaPrimeRénov’, les étapes absentes et les ventilations manquantes demeurent signalées comme telles. Les empreintes des fichiers sources ont été vérifiées. Cela ne constitue pas une nouvelle relecture humaine de chaque PDF ou une preuve que les sources publiques sont elles-mêmes exemptes d’erreurs.')
d.add_heading('Limites du crash test',1)
d.add_paragraph('Les contrôles chiffrés parcourent le jeu de données présent, mais les combinaisons de filtres et les interactions possibles sont trop nombreuses pour être toutes jouées dans un navigateur. Huit sélections variées ont fait l’objet du contrôle détaillé des affichages. La charge de quatre clients locaux ne mesure pas la capacité maximale du serveur de production. Aucun arrêt volontaire de Docker, du réseau ou de Windows n’a été provoqué.')
if fails:
 d.add_heading('Anomalies à examiner',1)
 for c in r['cases']:
  if c['status']!='pass':d.add_paragraph(c['id']+' : '+str(c['result'].get('error','Échec'))[:350])
d.add_heading('Relancer et consulter',1)
d.add_paragraph('Lanceur : tools\\LANCER-CRASH-TEST.cmd, dans D:\\ChatGPT\\docker\\budget. Il utilise Python et Chrome, sans appel à un modèle d’IA ni à RunPod. Docker doit être disponible et le site local doit répondre sur le port 8552.')
d.add_paragraph('Après une interruption, relancer le même lanceur conserve les scénarios terminés. Si le site, les données ou le script ont changé, utiliser un nouveau dossier avec --output ; les anciens résultats ne sont pas présentés comme une nouvelle certification.')
d.add_paragraph('Rapport détaillé : reports\\crash-test-site-20260924\\rapport.html. Résultats machine : report.json. Checkpoints : checkpoints.sqlite. Ces fichiers sont conservés avec le projet.')
for element in [d._element,d.styles.element]:
 for border in list(element.iter(qn('w:pBdr'))):border.getparent().remove(border)
d.save(out/'Nos Deniers Controle de solidite et des chiffres.docx')
print(out/'Nos Deniers Controle de solidite et des chiffres.docx')
