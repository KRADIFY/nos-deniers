
from pathlib import Path
import json
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'reports/finalisation-20260919'
audit=json.loads((OUT/'numeric-audit-after.json').read_text(encoding='utf8'))
doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=Cm(1.8);sec.left_margin=sec.right_margin=Cm(2)
styles=doc.styles
for name in ['Normal','Body Text','Title','Heading 1','Heading 2']:
 styles[name].font.name='Aptos';styles[name].font.color.rgb=RGBColor(0,0,0)
styles['Normal'].font.size=Pt(10.5);styles['Normal'].paragraph_format.space_after=Pt(7)
styles['Normal'].paragraph_format.line_spacing=1.08
styles['Title'].font.size=Pt(25);styles['Title'].paragraph_format.space_after=Pt(12)
styles['Heading 1'].font.size=Pt(16);styles['Heading 1'].paragraph_format.space_before=Pt(12)
styles['Heading 2'].font.size=Pt(12);styles['Heading 2'].paragraph_format.space_before=Pt(9)
header=sec.header.paragraphs[0];header.text='NOS DENIERS  ·  LexMachine';header.style='Caption'
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('19 septembre 2026  ·  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
def p(t,style=None):return doc.add_paragraph(t,style)
def h(t):doc.add_heading(t,1)
def link(label,url,paragraph=None):
 paragraph=paragraph or doc.add_paragraph()
 x=OxmlElement('w:hyperlink');x.set(qn('r:id'),paragraph.part.relate_to(url,RT.HYPERLINK,is_external=True))
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'173F5F');pr.append(c)
 u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u);r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);x.append(r);paragraph._p.append(x)
def table(headers,rows,widths=None):
 t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
 for cell,txt in zip(t.rows[0].cells,headers):
  cell.text=txt
  for r in cell.paragraphs[0].runs:r.bold=True
  shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EDF3F7');cell._tc.get_or_add_tcPr().append(shade)
 for row in rows:
  for cell,txt in zip(t.add_row().cells,row):cell.text=txt
 for row in t.rows:
  trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
  for i,cell in enumerate(row.cells):
   if widths:cell.width=Cm(widths[i])
   for par in cell.paragraphs:
    par.paragraph_format.space_after=Pt(5);par.paragraph_format.space_before=Pt(4)
    for r in par.runs:r.font.size=Pt(9.5)
 return t
def page():doc.add_page_break()

doc.add_paragraph('Nos Deniers',style='Title')
p('Bilan après vectorisation et préparation de la livraison',style='Subtitle')
p('La version locale permet désormais de chercher dans le corpus vectorisé et d’ouvrir les passages avec leurs références. Elle conserve les tableaux financiers, les exclusions, l’inflation et les exports. Les crédits proposés pour la mission Écologie en 2026 ont été ajoutés après vérification du PAP.')
p('Cette livraison constitue une avancée utilisable, mais pas une couverture complète de chaque poste du budget de l’État. Les manques sont détaillés ci-après. Le site public conserve sa version précédente jusqu’au transfert et à la publication de ce paquet.')
h('Ce qui est disponible')
table(['Fonction','État vérifié'],[
 ('Recherche documentaire','4 051 433 passages, issus de 4 439 références documentaires. Recherche par mots, par sens et reclassement des résultats.'),
 ('Preuves consultables','Passage complet, document source et page quand ils sont disponibles. L’année du filtre est celle du document.'),
 ('Tableaux budgétaires','120 610 observations structurées, dont 34 nouvelles pour le PAP Écologie 2026.'),
 ('Contrôles du site','167 tests Python réussis. Parcours navigateur vérifiés : recherche, exclusions, exports et affichage mobile.')
],[4.1,12.9])
h('Pour essayer avec votre fille')
p('Ouvrir Mission Écologie pour les comparaisons financières. Dans Rapports et documents, choisir Recherche sémantique, saisir par exemple « MaPrimeRénov crédits consommés », sélectionner l’année puis cliquer sur Rechercher. Ouvrir la page source avant de reprendre un chiffre.')
link('Ouvrir la version locale', 'http://127.0.0.1:8552/?scope=TA&start=2021&end=2026&measure=CP')
p('La recherche s’exécute sur le processeur, sans appel à RunPod. Les sept requêtes de contrôle ont pris environ 2 à 34 secondes, selon le démarrage du modèle et la question. Une seule recherche lourde s’exécute à la fois.')

page();h('Mission Écologie et fiabilité des chiffres')
table(['Périmètre','Couverture de cette version'],[
 ('2017 à 2025','Séries annuelles PLF, LFI et consommation disponibles selon les sources et le niveau sélectionné. Les trous restent visibles et les changements de nomenclature ne sont pas tous neutralisés.'),
 ('Actions et sous actions','Détails Écologie déjà intégrés, notamment 2023–2025, avec sources. Le détail n’est pas exhaustif pour toutes les années ni toutes les missions.'),
 ('Réserves Écologie 2017 à 2025','76 tableaux intégrés. Réserve initiale, mouvements et solde restent distincts des crédits consommés. Les tableaux absents ne deviennent pas des zéros.'),
 ('2026','LFI présente ; PLF et prévisions FdC/AdP ajoutés pour Écologie. Le PLF national reste partiel. Pas de consommation annuelle 2026 complète.')
],[4.2,12.8])
doc.add_heading('Ajout contrôlé dans le PAP 2026',2)
table(['Montants en euros courants','AE','CP'],[
 ('Crédits proposés en PLF','24 237 621 537','21 814 445 422'),
 ('FdC et AdP prévus','3 525 099 960','3 573 802 460')
],[7.4,4.8,4.8])
p('Source : PAP Écologie 2026, pages PDF 17 à 20. Les lignes PLF sont distinguées des lignes LFI 2025 du même tableau. Dix programmes sont chiffrés ; sept ont une prévision FdC/AdP renseignée. Les autres cases blanches restent absentes.')
link('Consulter le PAP Écologie 2026 à la page 17','http://127.0.0.1:8552/api/download/b3b6f06b6363f6f7df96#page=17')
doc.add_heading('Ce que couvrent les vérifications',2)
p(f"Intégrité SQLite conforme ; 120 576 observations antérieures conservées à l’identique ; 74 fichiers sources utilisés par les données structurées vérifiés par empreinte ; {audit['page_references_checked']} références de pages contrôlées ; 264 rapprochements de totaux de missions en 2024–2025. Les montants du nouvel ajout ont aussi été relus visuellement dans le PDF.")
p('Ces contrôles ne certifient pas chaque nombre du corpus documentaire. Les 97 signalements historiques de rapprochement des crédits ouverts restent tracés ; ils ne désignent pas 97 PDF corrompus. Les valeurs sources n’ont pas été artificiellement ajustées pour faire disparaître les écarts.')

page();h('MaPrimeRénov et données restant à isoler')
p('Le périmètre suivi est celui des crédits de l’État identifiés pour MaPrimeRénov’. Les aides versées par l’Anah aux bénéficiaires, les crédits transférés par l’État et les autres financements de la rénovation ne sont pas interchangeables.')
table(['Cas','Utilisable ou restant à établir'],[
 ('Consommation nationale 2021–2024','Montants déjà intégrés et sourcés. Leur précision varie selon les publications.'),
 ('Écologie 2020','LFI et consommation du programme 174 documentées. Cela ne constitue pas à lui seul une série nationale complète de toutes les étapes.'),
 ('Écologie 2025','Le transfert vers le programme 135 est documenté. L’exclusion dans Écologie ne signifie pas que MaPrimeRénov’ vaut zéro au niveau national.'),
 ('National 2020 et 2025–2026','Montants exacts à isoler pour les étapes manquantes. Les résultats de recherche apportent des pièces, sans lever toutes les ambiguïtés.'),
 ('PLF et mouvements spécifiques','PLF, certaines LFI, crédits ouverts, réserves et mouvements propres à MaPrimeRénov’ restent à compléter selon l’année.')
],[5.1,11.9])
doc.add_heading('Deux limites concrètes retrouvées dans les sources',2)
p('Le RAP Cohésion des territoires 2025 présente des versements d’intervention à l’Anah couvrant notamment MaPrimeRénov’, MaPrimeAdapt’ et la lutte contre l’habitat indigne. Leur total ne peut donc pas être retiré comme s’il appartenait intégralement à MaPrimeRénov’. Voir pages PDF 122 et 142.')
link('Ouvrir le RAP Cohésion des territoires 2025','http://127.0.0.1:8552/api/download/d2c406b7e2bc73904e50#page=122')
p('Le jaune budgétaire consacré à la rénovation énergétique, annexé au PLF 2026, utilise une estimation de la part des aides de l’Anah consacrée à la rénovation énergétique. Elle ne fournit pas une ventilation comptable exacte de MaPrimeRénov’. Voir pages PDF 11 et 17.')
link('Ouvrir le jaune Rénovation énergétique 2026','http://127.0.0.1:8552/api/download/dbfe025ccdf7c863dd3d#page=17')
p('Pour compléter ces cases, il faut une ventilation AE/CP par dispositif, année et étape budgétaire, puis rapprocher cette ventilation des totaux du programme et de l’opérateur. Aucun pourcentage arbitraire du programme ou du budget de l’Anah n’a été appliqué.')

page();h('Livraison et suite du développement')
p('Le paquet est préparé localement pour budget.lexmachine.net. Il contient le site, le moteur de recherche, les changements de la base financière et un vérificateur conservant les références de la version publique précédente. La mise en ligne n’a pas été exécutée pendant cette étape.')
table(['Élément','Organisation retenue'],[
 ('Données déjà publiques','Les documents et les données inchangés sont réutilisés sur le serveur. L’ancienne version reste disponible pour revenir en arrière.'),
 ('Index de recherche','Environ 30,1 Go : texte, index lexical, vecteurs quantifiés et références. Les sorties RunPod d’origine restent conservées sur F.'),
 ('Moteur Docker','Archive compressée d’environ 3,2 Go. Aucun téléchargement de modèle au démarrage du serveur.'),
 ('Transfert','Plan local et script SFTP reprenable. Chaque fichier reçu est contrôlé par empreinte avant publication. Environ 33,3 Go au total.'),
 ('Exploitation','Moteur privé, sans port public, accès aux données en lecture seule. Limite de mémoire de 2,2 Gio pour la recherche.')
],[4.6,12.4])
doc.add_heading('Ce qui reste dans le plan initial',2)
p('Compléter les montants MaPrimeRénov’ cités page précédente ; extraire et contrôler davantage de réserves et mouvements hors Écologie ; étendre les PLF 2026 et les détails par action aux autres missions ; poursuivre les correspondances de nomenclature entre années.')
p('L’analyse rédigée automatiquement par une IA à partir d’une sélection figée, la recherche fédérée en direct dans toutes les connexions externes et le suivi complet de la navette parlementaire ne sont pas livrés dans cette version. La recherche locale constitue désormais une base pour ces étapes.')
p('Une absence de résultat de recherche ne prouve pas qu’une donnée n’existe pas. Certaines références ne disposent que d’un texte extrait ou d’un lien externe ; la possibilité d’ouvrir le fichier est indiquée dans le résultat.')
doc.add_heading('Emplacements à conserver',2)
link('Dossier de livraison et instructions', (ROOT/'deploy/update-20260919-recherche').as_uri())
link('Dossier des contrôles et du bilan', OUT.as_uri())
link('Index documentaire utilisé par le site', Path('D:/LexMachine/NosDeniers/search_20260919').as_uri())
p('Aucune nouvelle vectorisation ni dépense RunPod n’a été engagée pour cette livraison. La génération source sur F a été conservée.')
out=OUT/'Nos Deniers - Bilan apres vectorisation.docx'
for part in [doc.styles.element,doc._element]:
 for border in list(part.xpath('//w:pBdr')):border.getparent().remove(border)
doc.save(out);print(out)
