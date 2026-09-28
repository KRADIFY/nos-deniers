"""Assemble and extract only additional documents, preserving source tiers and originals."""
from pathlib import Path
import sys,json,hashlib,shutil,zipfile,re,sqlite3,csv,io,datetime
from html.parser import HTMLParser
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'.runtime/python-libs'))
from consolidate_vectorisation_20260924 import ROOT,OUT,CONTROL,sha,save

TIERS={'official':'Source officielle. Extraction documentaire brute, non assimilable à un fait budgétaire validé.',
       'working':'Tableau de travail fourni par l’utilisateur. Compilation secondaire, à vérifier dans les sources officielles.',
       'internal':'Mémo interne Nos Deniers au 24 septembre 2026. État de travail historique, non source officielle.'}

def walk(v):
    if isinstance(v,dict):
        yield v
        for x in v.values():yield from walk(x)
    elif isinstance(v,list):
        for x in v:yield from walk(x)

def cleanname(s):
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]','_',s).strip(' .')[:150]

def copy_checked(src,dest):
    dest.parent.mkdir(parents=True,exist_ok=True)
    if dest.exists():assert sha(src)==sha(dest),(src,dest)
    else:shutil.copy2(src,dest)
    assert sha(src)==sha(dest)

def assemble():
    inv=json.loads((CONTROL/'inventaire-initial.json').read_text('utf-8'))
    meta={}
    for e in inv['candidates']:
        for m in e['metadata']:meta.setdefault(e['sha256'],[]).append(m)
    for p in (ROOT/'reports').rglob('*.json'):
        if not re.search(r'202609(19|20|21|22|23|24)',str(p)) or p.stat().st_size>3_000_000:continue
        if any(s in str(p).lower() for s in ('vectorisation','page_inventory','pages.json','before','candidate.sqlite')):continue
        try:d=json.loads(p.read_text('utf-8-sig'))
        except (ValueError,OSError):continue
        for m in walk(d):
            h=m.get('sha256') or m.get('source_sha256')
            if isinstance(h,str) and len(h)==64 and any(m.get(k) for k in ('url','source_url','final_url')):meta.setdefault(h,[]).append(m)
    records=[]; decisions=[]
    for e in inv['candidates']:
        paths=[Path(p) for p in e['paths']];p=paths[0];suffix=p.suffix.lower()
        supplied=any('Desktop\\Nouveau dossier' in str(q) for q in paths)
        tier='official';action='index';reason='nouvelle source documentaire'
        if e['indexed']:action='reuse';reason='empreinte présente dans les passages de l’index actif'
        elif suffix=='.zip':action='archive';reason='conteneur conservé ; indexer les membres utiles une seule fois'
        elif supplied and suffix=='.xlsx':tier='working'
        elif supplied and suffix=='.xls':action='archive';reason='copie utilisateur ; mêmes valeurs que le classeur officiel 2023 archivé, contrôle à refaire'
        elif suffix in ('.docx','.doc') and 'manual-documents' not in str(p):
            tier='internal'
            if supplied and ('actualisees' in p.name or 'Méta-audit' in p.name or 'Lecture des onglets' in p.name):action='context_only';reason='mémo interne ; conserver hors recherche budgétaire par défaut'
            elif supplied:action='archive';reason='ancienne liste remplacée par la liste actualisée'
            else:action='exclude';reason='ancien audit ou sortie de contrôle déjà conservée dans le projet'
        elif suffix=='.xlsx':action='exclude';reason='export du site pour contrôle, non nouvelle source'
        elif suffix=='.pdf' and ('Nos Deniers' in p.name or p.name=='memo.pdf'):action='exclude';reason='rendu ou ancien audit interne'
        elif suffix in ('.html','.htm') and ('before' in str(p) or 'explorer' in p.name):action='exclude';reason='copie de l’interface, non source documentaire'
        decisions.append(dict(sha256=e['sha256'],paths=e['paths'],action=action,reason=reason,tier=tier))
        if action not in ('index','context_only','archive'):continue
        folder={'official':'01_sources_officielles','working':'02_tableaux_de_travail','internal':'03_memos_internes'}[tier]
        if action=='archive':folder='04_archives_non_a_vectoriser'
        target=OUT/folder/(e['sha256'][:12]+'-'+cleanname(p.name))
        copy_checked(p,target)
        ms=meta.get(e['sha256'],[])
        m=next((m for m in ms if m.get('url') or m.get('source_url')),ms[0] if ms else {})
        url=m.get('url') or m.get('source_url') or m.get('final_url') or ''
        if not url and 'annexes-circulaire' in str(p):url='https://www.budget.gouv.fr/documentation/fid-download/77566'
        if not url and p.name.startswith(('d641','d762','73c9')):
            url='https://www.budget.gouv.fr/documentation/documents-budgetaires/exercice-'+('2023' if p.name.startswith('d641') else '2024' if p.name.startswith('d762') else '2025')+'/plrg-'+('2023' if p.name.startswith('d641') else '2024' if p.name.startswith('d762') else '2025')
        title=m.get('title') or p.stem
        rec=dict(id=e['sha256'][:20],sha256=e['sha256'],bytes=e['bytes'],title=title,path=target.relative_to(OUT).as_posix(),original_paths=e['paths'],url=url,source_tier=tier,format=suffix[1:],stage='documentation' if tier=='official' else tier,years_title=m.get('years_title') or re.findall(r'20(?:1\d|2\d)',p.name),action=action,numeric_status='raw_not_validated_facts',authority_note=TIERS[tier])
        records.append(rec)
    # Archive members: safe, deterministic filenames, exact uncompressed bytes and parent provenance.
    for archive in list(records):
        if archive['format']!='zip':continue
        with zipfile.ZipFile(OUT/archive['path']) as z:
            assert z.testzip() is None
            for n in z.namelist():
                if n.endswith('/'):continue
                b=z.read(n);h=hashlib.sha256(b).hexdigest()
                if any(x['sha256']==h for x in records):continue
                ext=Path(n).suffix.lower()
                if ext not in ('.csv','.pdf','.doc','.docx'):raise ValueError('Unexpected archive member '+n)
                target=OUT/'01_sources_officielles'/('plrg2024' if 'PLRG2024' in archive['path'] else 'annexes-rap2024')/(h[:12]+'-'+cleanname(Path(n).name))
                target.parent.mkdir(parents=True,exist_ok=True)
                if target.exists():assert sha(target)==h
                else:target.write_bytes(b)
                records.append(dict(id=h[:20],sha256=h,bytes=len(b),title=Path(n).stem,path=target.relative_to(OUT).as_posix(),original_paths=[],archive_sha256=archive['sha256'],archive_member=n,url=archive['url'] or ('https://www.budget.gouv.fr/documentation/fid-download/78647' if 'PLRG2024' in archive['path'] else 'https://www.budget.gouv.fr/documentation/fid-download/77566'),source_tier='official',format=ext[1:],stage='plrg' if 'PLRG2024' in archive['path'] else 'methodologie',years_title=['2024'],action='index',numeric_status='raw_not_validated_facts',authority_note=TIERS['official']))
    # Bibliographic URLs known from the prior receipt; never invent a direct download URL.
    for r in records:
        if r['title']=='PLRG2024-donnees-chiffrees-standard-ouvert':r['url']='https://www.budget.gouv.fr/documentation/fid-download/78647'
    save(CONTROL/'sources.json',records);save(CONTROL/'decisions-doublons.json',decisions)
    print(json.dumps(dict(assembled=len(records),index=sum(r['action']=='index' for r in records),internal=sum(r['action']=='context_only' for r in records),archive=sum(r['action']=='archive' for r in records)),ensure_ascii=False))

class ExtractHTML(HTMLParser):
    def __init__(self):super().__init__(convert_charrefs=True);self.depth=0;self.parts=[];self.anchor='';self.heading='';self.anchors=[]
    def handle_starttag(self,tag,attrs):
        if tag in ('script','style','nav','header','footer'):self.depth+=1
        if self.depth:return
        a=dict(attrs);ident=a.get('id') or a.get('name')
        if ident:self.anchor=ident;self.anchors.append(dict(anchor=ident,offset=len(''.join(self.parts))))
        if tag in ('p','div','tr','table','h1','h2','h3','h4','li','br'):self.parts.append('\n')
        elif tag in ('td','th'):self.parts.append(' | ')
    def handle_endtag(self,tag):
        if tag in ('script','style','nav','header','footer') and self.depth:self.depth-=1
        if not self.depth and tag in ('p','tr','h1','h2','h3','h4','li'):self.parts.append('\n')
    def handle_data(self,data):
        if not self.depth:self.parts.append(data)

def decode(b):
    for enc in ('utf-8-sig','cp1252'):
        try:return b.decode(enc),enc
        except UnicodeDecodeError:pass
    raise ValueError('Unknown text encoding')

def extract():
    from pypdf import PdfReader
    from docx import Document
    import openpyxl,xlrd
    recs=json.loads((CONTROL/'sources.json').read_text('utf-8'))
    output=OUT/'05_extractions';output.mkdir(exist_ok=True)
    dbpath=output/'tableaux.sqlite'
    if dbpath.exists():dbpath.unlink()  # Only this generated file, never an active database.
    db=sqlite3.connect(dbpath)
    db.executescript('create table sheets(source_id text,sheet text,rows integer,columns integer,merged_json text,primary key(source_id,sheet)); create table cells(source_id text,sheet text,row integer,col integer,address text,value_json text,formula text,hyperlink text,number_format text,primary key(source_id,sheet,row,col)); create table csv_rows(source_id text,row integer,fields_json text,primary key(source_id,row)); create table csv_headers(source_id text primary key,headers_json text,encoding text,delimiter text);')
    segments=[];report=[];pending=[]
    def add(r,loc,text,**more):
        if not text.strip():return
        segments.append(dict(source_id=r['id'],source_sha256=r['sha256'],format=r['format'],locator=loc,text=text,source_tier=r['source_tier'],**more))
    for r in recs:
        if r['action'] not in ('index','context_only'):continue
        p=OUT/r['path'];assert sha(p)==r['sha256'];before=len(segments);fmt=r['format'];item={'source_id':r['id'],'path':r['path'],'format':fmt}
        try:
            if fmt=='pdf':
                reader=PdfReader(p);item['pages']=len(reader.pages);blank=[]
                for i,page in enumerate(reader.pages,1):
                    t=page.extract_text(extraction_mode='layout') or ''
                    if len(t.strip())<35:
                        if len(page.images)>0:pending.append(dict(kind='ocr',source_id=r['id'],page=i,path=r['path']));continue
                        blank.append(i)
                    add(r,f'page:{i}',t,extraction_method='pypdf_layout',table_layout='raw_unvalidated')
                item['blank_pages']=blank
            elif fmt in ('html','htm'):
                text,enc=decode(p.read_bytes());hp=ExtractHTML();hp.feed(text);text=''.join(hp.parts)
                for start in range(0,len(text),12000):
                    eligible=[a for a in hp.anchors if a['offset']<=start];anchor=eligible[-1]['anchor'] if eligible else ''
                    add(r,f'html:#{anchor}/chars:{start}-{min(start+12000,len(text))}',text[start:start+12000],anchors=[],extraction_method='html_parser_preserving_table_cells')
                item['encoding']=enc
            elif fmt=='doc':
                converted=output/'documents-convertis'/(r['id']+'.docx')
                if not converted.exists():pending.append(dict(kind='doc_conversion',source_id=r['id'],path=r['path']));continue
                doc=Document(converted)
                for i,p0 in enumerate(doc.paragraphs,1):add(r,f'paragraph:{i}',p0.text,extraction_method='word_conversion')
                for i,table in enumerate(doc.tables,1):
                    for j,row in enumerate(table.rows,1):add(r,f'table:{i}/row:{j}',' | '.join(c.text for c in row.cells),extraction_method='word_conversion')
            elif fmt=='docx':
                doc=Document(p)
                for i,p0 in enumerate(doc.paragraphs,1):add(r,f'paragraph:{i}',p0.text,extraction_method='ooxml')
                for i,table in enumerate(doc.tables,1):
                    for j,row in enumerate(table.rows,1):add(r,f'table:{i}/row:{j}',' | '.join(c.text for c in row.cells),extraction_method='ooxml')
            elif fmt in ('xlsx','xls'):
                if fmt=='xls':
                    book=xlrd.open_workbook(p,formatting_info=True);sheets=[(s.name,s.nrows,s.ncols,[[c.value for c in row] for row in s.get_rows()],s.merged_cells,s) for s in book.sheets()]
                else:
                    book=openpyxl.load_workbook(p,data_only=True,read_only=False);form=openpyxl.load_workbook(p,data_only=False,read_only=False)
                    sheets=[(s.title,s.max_row,s.max_column,[[c.value for c in row] for row in s],list(map(str,s.merged_cells.ranges)),s) for s in book]
                total=0
                for name,nrows,ncols,rows,merged,sheet in sheets:
                    db.execute('insert into sheets values(?,?,?,?,?)',(r['id'],name,nrows,ncols,json.dumps(merged)))
                    head=[]
                    for i,row in enumerate(rows,1):
                        cells=[]
                        for j,val in enumerate(row,1):
                            formula=link=nformat=''
                            if fmt=='xlsx':
                                c=sheet.cell(i,j);f=form[name].cell(i,j)
                                formula=f.value if f.data_type=='f' else '';link=c.hyperlink.target if c.hyperlink else '';nformat=c.number_format
                            if val in (None,'') and not formula:continue
                            addr=openpyxl.utils.get_column_letter(j)+str(i)
                            db.execute('insert into cells values(?,?,?,?,?,?,?,?,?)',(r['id'],name,i,j,addr,json.dumps(val,ensure_ascii=False,default=str),formula,link,nformat))
                            total+=1;cells.append(f'{addr}={val if val is not None else "[valeur calculée non disponible]"}'+(f' [formule {formula}]' if formula else ''))
                        if not cells:continue
                        if i<=8:head.extend(cells)
                        add(r,f'sheet:{name}/row:{i}',f'Onglet {name}. Ligne {i}. '+' ; '.join(cells),extraction_method='cell_coordinates',header_context=' ; '.join(head),table_layout='coordinates_preserved')
                item['cells']=total;item['sheets']=len(sheets)
                if fmt=='xlsx':book.close();form.close()
            elif fmt=='csv':
                text,enc=decode(p.read_bytes());text=text.replace('\r\r\n','\n')
                sample=text[:12000];dialect=csv.Sniffer().sniff(sample,delimiters=';,\t');rows=list(csv.reader(io.StringIO(text,newline=None),dialect));rows=[row for row in rows if any(v.strip() for v in row)]
                headers=rows[0];db.execute('insert into csv_headers values(?,?,?,?)',(r['id'],json.dumps(headers,ensure_ascii=False),enc,dialect.delimiter))
                for i,row in enumerate(rows[1:],2):
                    if len(row)!=len(headers):raise ValueError(f'CSV width mismatch row {i}: {len(row)}/{len(headers)}')
                    db.execute('insert into csv_rows values(?,?,?)',(r['id'],i,json.dumps(row,ensure_ascii=False)))
                    content=' ; '.join(f'{h.strip() or "colonne "+str(k)}={v if v.strip() else "[vide]"}' for k,(h,v) in enumerate(zip(headers,row),1))
                    add(r,f'csv:row:{i}',content,extraction_method='csv_fields',table_layout='headers_preserved')
                item.update(rows=len(rows)-1,columns=len(headers),encoding=enc,delimiter=dialect.delimiter)
            else:raise ValueError('Unsupported format '+fmt)
        except Exception as e:
            pending.append(dict(kind='extraction_error',source_id=r['id'],path=r['path'],error=str(e)))
        item['segments']=len(segments)-before;report.append(item)
        print(r['id'],fmt,item['segments'],flush=True)
    # OCR transcriptions stay distinct from original text and never certify monetary amounts.
    ocrfile=output/'ocr-pages.json'
    if ocrfile.exists():
        ocr=json.loads(ocrfile.read_text('utf-8-sig'));byid={r['id']:r for r in recs};done=set()
        for page in ocr:
            if page.get('text','').strip():
                add(byid[page['source_id']],f'page:{page["page"]}',page['text'],extraction_method='windows_ocr_fr',table_layout='ocr_unvalidated');done.add((page['source_id'],page['page']))
        pending=[x for x in pending if x['kind']!='ocr' or (x['source_id'],x['page']) not in done]
    db.commit();assert db.execute('pragma integrity_check').fetchone()[0]=='ok';db.close()
    save(CONTROL/'extraction-report.json',report);save(CONTROL/'preparation-pending.json',pending)
    with (output/'segments.jsonl').open('w',encoding='utf-8') as f:
        for s in segments:f.write(json.dumps(s,ensure_ascii=False,default=str)+'\n')
    save(CONTROL/'extraction-summary.json',dict(sources=len(report),segments=len(segments),pending=len(pending),raw_amounts_certified=False,main_index_modified=False))
    print('SUMMARY',len(report),len(segments),'pending',len(pending),flush=True)

if __name__=='__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    if '--assemble' in sys.argv:assemble()
    if '--extract' in sys.argv:extract()
