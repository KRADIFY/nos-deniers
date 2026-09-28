from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

OUT=Path(__file__).resolve().parents[1]/'reports/memo-sites'
OUT.mkdir(parents=True,exist_ok=True)
doc=Document()
sec=doc.sections[0]
sec.page_height=Cm(29.7); sec.page_width=Cm(21)
sec.top_margin=Cm(1.7); sec.bottom_margin=Cm(1.7)
sec.left_margin=Cm(2); sec.right_margin=Cm(2)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    st=doc.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
normal=doc.styles['Normal']; normal.font.size=Pt(10.5)
normal.paragraph_format.space_after=Pt(5)
normal.paragraph_format.line_spacing=1.05
doc.styles['Title'].font.size=Pt(23)
doc.styles['Heading 1'].font.size=Pt(17)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(10)

def link(p,label,url):
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r'); pr=OxmlElement('w:rPr')
    c=OxmlElement('w:color'); c.set(qn('w:val'),'175785'); pr.append(c)
    u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u)
    r.append(pr); t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

def entry(name,rank,body,fresh,links):
    doc.add_heading(name,2)
    p=doc.add_paragraph();p.add_run(rank+' • ').bold=True;p.add_run(body)
    p=doc.add_paragraph();p.add_run('Actualité : ').bold=True;p.add_run(fresh)
    p=doc.add_paragraph()
    for i,(label,url) in enumerate(links):
        if i:p.add_run('  |  ')
        link(p,label,url)

doc.add_paragraph('Les sites de suivi budgétaire',style='Title')
doc.add_paragraph('Mémo de visite pour Nos Deniers • 9 septembre 2026')
doc.add_paragraph('Pour regarder les interfaces, commencer par Les fonds du sujet, Polidata et France Transparence. Pour construire le service, les sources officielles et Data-État sont les pistes prioritaires. Aucun outil étudié ne démontre à lui seul tout le suivi PLF, LFI, consommation, réserve et mouvements au niveau de chaque action.')
doc.add_heading('Les quatre premières visites',1)
entry('Les fonds du sujet','Fort pour la lisibilité',
    'Site indépendant gratuit. Le benchmark relève une navigation 2017–2026 par missions et programmes, surtout autour des crédits votés. À regarder : choix des années, courbes et lecture des enveloppes. Le cycle complet de gestion reste à établir.',
    'Activité éditoriale attestée en septembre 2026. Cela ne certifie pas la fraîcheur de chaque tableau budgétaire.',
    [('Budget de l’État','https://lesfondsdusujet.info/etat'),('Articles récents','https://lesfondsdusujet.info/blog')])
entry('Polidata','Ciblé pour la présentation',
    'Comparateur public plus léger, décrit dans le benchmark à partir du budget vert : exécution 2024, LFI 2025, PLF 2026 en CP. Utile pour présenter trois situations clairement ; moins adapté à dix années et à la chronologie des gels.',
    'Millésime 2026 relevé dans le benchmark du 8 septembre. Page non relue avec succès le 9 septembre ; date de lancement et maintenance non établies.',
    [('Comparateur PLF 2026','https://polidata.fr/budget/plf2026')])
entry('France Transparence','Fort pour les sources et les rapprochements',
    'Tableau de bord public des finances et de la vie politique. Le code documente des imports de sources officielles, une base SQLite et l’affichage de la fraîcheur des données. Intéressant pour la traçabilité ; couverture complète par sous-action non démontrée.',
    'Le dépôt décrit une construction en août 2026 et une reconstruction quotidienne annoncée. La fréquence annoncée ne prouve pas que chaque source se met à jour tous les jours.',
    [('Dépenses','https://francetransparence.fr/depenses/'),('Code et documentation','https://github.com/koutakou/france-transparence')])
entry('Budget Data État','Très fort pour le fonctionnement métier',
    'Service administratif de recherche de financements, engagements et paiements, avec tableaux et filtres. Regarder les tutoriels pour comprendre les usages. Code MIT identifié dans le benchmark ; accès aux données réservé aux agents autorisés.',
    'Fiche officielle en accélération et tutoriels disponibles. Activité des dépôts relevée au 13 août 2026 dans le benchmark ; cette date n’a pas été reconfirmée ici.',
    [('Tutoriels','https://www.dataregion.fr/tutoriels'),('Projet officiel','https://beta.gouv.fr/startups/data.etat.html'),('Code','https://github.com/dataregion/DataEtatRegion-back')])

doc.add_page_break();doc.add_heading('Les sources et outils parlementaires',1)
entry('Budget gouv fr','Indispensable pour les documents officiels',
    'Catalogue des lois de finances, PAP, RAP, nomenclatures et annexes. Fournit les preuves nécessaires aux tableaux de Nos Deniers. L’interface documentaire impose de parcourir dossiers et sous-dossiers ; c’est la collecte actuellement automatisée.',
    'Catalogue officiel consultable, publications par exercice dont 2026. Un PAP 2026 et un RAP 2025 décrivent des étapes et années différentes.',
    [('Documents budgétaires','https://www.budget.gouv.fr/documentation/documents-budgetaires'),('Exploration par mission','https://www.budget.gouv.fr/budget-etat/mission')])
entry('Data économie et data gouv','Indispensable pour les fichiers structurés',
    'Portails publics complémentaires au site Budget : jeux CSV, classeurs, métadonnées et API selon les jeux. Ce sont les meilleures entrées pour alimenter les calculs. La finesse et la couverture changent selon l’année ; certaines colonnes exigent un rapprochement avec les documents originaux.',
    'Jeux récents PLRG 2024–2025 vérifiés dans le benchmark du 8 septembre. Il faut consulter la date de chaque ressource, pas seulement celle du catalogue.',
    [('Catalogue économique','https://data.economie.gouv.fr/explore/'),('Catalogue national','https://www.data.gouv.fr/'),('Exemple PLRG 2024','https://data.economie.gouv.fr/explore/dataset/plrg-2024/')])
entry('Legiwatch','Fort pour la navette parlementaire',
    'Service commercial de veille, recherche, alertes et suivi des amendements. Le benchmark identifie un suivi des crédits pendant le PLF. À regarder pour les parcours quotidiens d’un assistant parlementaire ; historique de consommation et réserve non démontrés.',
    'Offre commerciale et cas d’usage consultables le 9 septembre 2026. Fonctionnalités décrites par l’éditeur, sans test de compte client ; prix et intégration à confirmer avant achat.',
    [('Présentation','https://www.legiwatch.fr/'),('Cas d’usage PLF','https://www.legiwatch.fr/cas-d-usage'),('Module du benchmark','https://www.legiwatch.fr/plateforme/suivi-plf/')])
entry('Contexte','Fort pour comprendre les arbitrages',
    'Média professionnel sur abonnement : analyses, documents de préparation, alertes, suivi des articles et amendements, exports Excel annoncés. Apporte une explication politique et documentaire ; une base exhaustive de montants annuels n’est pas établie.',
    'Média fondé en 2013. Offre dédiée au budget 2026 relue le 9 septembre : ce service ancien reste pertinent pour la période actuelle.',
    [('Offre budget 2026','https://about.contexte.com/budget-plf-plfss')])
doc.add_paragraph('À retenir : une consultation gratuite, un abonnement et un code ouvert donnent des possibilités différentes. Une éventuelle intégration dans Nos Deniers devra être qualifiée séparément de la simple visite du site.')

doc.add_page_break();doc.add_heading('Les compléments et références de conception',1)
entry('Dépenses éclairées','Prometteur pour comprendre la dépense',
    'Projet AIFE avec soutien DINUM : enrichir les dépenses Chorus grâce aux pièces de commande, notamment pour mieux identifier l’objet et les prestataires. Peut éclairer « à quoi cela a servi ». Ne constitue pas une base publique exhaustive disponible pour Nos Deniers.',
    'Fiche officielle marquée en construction lors de la consultation du 9 septembre. Service à suivre ; accès exploitable à confirmer.',
    [('Fiche officielle','https://beta.gouv.fr/startups/depenses-eclairees.html')])
entry('Gobierto Presupuestos','Fort pour les idées d’interface',
    'Solution espagnole commerciale : exploration progressive des postes, évolution annuelle, exécution et comparaison de versions. Bon exemple pour articuler vue générale et détail. Adaptation aux sources françaises et à la LOLF nécessaire.',
    'Offre et démonstration proposées sur la page relue le 9 septembre. Code public identifié dans le benchmark ; maintenance du composant précis à vérifier avant reprise.',
    [('Présentation et démonstration','https://www.gobierto.es/transparencia/presupuestos'),('Code','https://github.com/PopulateTools/gobierto')])
entry('USAspending','Fort pour la traçabilité des financements',
    'Service public américain avec recherche des financements, comptes, bénéficiaires et API. À regarder pour les filtres et les liens entre montants et bénéficiaires. Référence de conception, sans données budgétaires françaises.',
    'Documentation API accessible le 9 septembre. Les dates de mise à jour doivent se lire dans les jeux américains concernés ; elles ne renseignent pas la couverture française.',
    [('Site public','https://www.usaspending.gov/'),('Documentation API','https://api.usaspending.gov/docs/endpoints')])
entry('Manty Budget','Secondaire pour notre périmètre',
    'Logiciel commercial de préparation budgétaire, notamment pour les collectivités : saisie par les services, arbitrages et connexion à leur logiciel financier. Utile pour étudier un espace de travail collaboratif ; aucune base nationale LOLF prête à consulter n’est démontrée.',
    'Page produit et proposition de démonstration accessibles le 9 septembre. L’éditeur annonce plus de 250 collectivités utilisatrices ; chiffre déclaratif, non audité.',
    [('Produit et démonstration','https://www.manty.eu/budget')])
doc.add_paragraph('Pour Nos Deniers, la priorité reste de relier montants officiels, postes historiques et événements de gestion. Les exemples étrangers ou territoriaux peuvent aider la présentation sans fournir cette base française.')

doc.add_page_break();doc.add_heading('Les pistes spécialisées et historiques',1)
entry('LexImpact','Complément pour simuler les réformes',
    'Service de l’Assemblée nationale pour évaluer des modifications fiscales et sociales. Très pertinent pour le travail parlementaire, mais finalité distincte du suivi des crédits votés puis consommés.',
    'Lancé en 2019 et pérennisé à l’Assemblée depuis 2020, selon sa fiche officielle. Celle-ci indique la poursuite de son amélioration ; ce n’est pas un simple prototype abandonné.',
    [('Simulateurs','https://leximpact.an.fr/'),('Historique officiel','https://beta.gouv.fr/startups/leximpact.html')])
entry('Comptes de l’État','Complément pour la vision patrimoniale',
    'Datavisualisation officielle du bilan et du résultat de l’État. Utile pour distinguer actifs, passifs et charges. La comptabilité générale complète les crédits budgétaires mais ne doit pas être additionnée à ceux-ci.',
    'Page officielle accessible le 9 septembre. Série 2015–2024 relevée dans le benchmark ; extension ultérieure non établie dans ce mémo.',
    [('Comptes de l’État','https://www.budget.gouv.fr/reperes/comptes_etat')])
entry('OpenSpending','Secondaire pour une reprise logicielle',
    'Projet international de publication et exploration de données financières. Les anciens composants ne sont pas une application française prête à installer. Une reprise exigerait une évaluation technique spécifique.',
    'Le benchmark relève des dépôts historiques anciens. Datopian documente aussi une refonte ultérieure : il serait donc excessif de qualifier tout OpenSpending d’abandonné.',
    [('Dépôt historique','https://github.com/openspending/openspending'),('Refonte par Datopian','https://www.datopian.com/blog/the-open-spending-revamp')])
entry('OpenBudget fr','Faible pour une intégration actuelle',
    'Ancienne visualisation citoyenne référencée sur data.gouv.fr autour du budget 2014. Peut donner des idées pédagogiques ; aucune continuité de couverture récente n’est démontrée.',
    'Référence historique. Une fiche de catalogue encore accessible ne prouve ni le fonctionnement de l’application d’origine ni l’actualisation des chiffres.',
    [('Fiche et lien d’origine','https://www.data.gouv.fr/reuses/openbudget-fr-le-budget-un-outil-de-la-democratie')])
doc.add_heading('Comment lire ce mémo',2)
doc.add_paragraph('Les appréciations de pertinence concernent le besoin de Nos Deniers. Les liens de chaque fiche servent à visiter le service et à retrouver les éléments décrits. Ce mémo actualise le benchmark du 8 septembre avec les consultations du 9 septembre 2026 ; les informations reprises du benchmark sont signalées lorsque la fraîcheur n’a pas été reconfirmée. Les espaces protégés et les offres payantes n’ont pas été testés.')
doc.core_properties.title='Les sites de suivi budgétaire pour Nos Deniers'
doc.core_properties.author='LexMachine'
path=OUT/'Memo des sites budgetaires Nos Deniers.docx'
for root in [doc.styles.element,doc.element]:
    for border in list(root.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.save(path)
print(path)

