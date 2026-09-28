"""Render the reviewed audit Markdown into an editable, linked Word report."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

OUT = Path(__file__).resolve().parents[1] / 'reports/audit-20260909'
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = section.bottom_margin = Inches(.7)
section.left_margin = section.right_margin = Inches(.8)
section.header_distance = section.footer_distance = Inches(.3)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    for elem in list(style.element.iter(qn('w:color'))):
        for attr in ['themeColor', 'themeTint', 'themeShade']:
            elem.attrib.pop(qn('w:' + attr), None)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing = 1.07
doc.styles['Title'].font.size = Pt(23)
doc.styles['Title'].paragraph_format.space_after = Pt(12)
for name, size in [('Heading 1', 16), ('Heading 2', 12.5), ('Heading 3', 11.5)]:
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].paragraph_format.space_before = Pt(14)
    doc.styles[name].paragraph_format.space_after = Pt(6)
    doc.styles[name].paragraph_format.keep_with_next = True

def hyperlink(paragraph, text, url):
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), paragraph.part.relate_to(url.strip('<>'), RT.HYPERLINK, is_external=True))
    r = OxmlElement('w:r'); props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '175785'); props.append(color)
    under = OxmlElement('w:u'); under.set(qn('w:val'), 'single'); props.append(under)
    r.append(props); text_node = OxmlElement('w:t'); text_node.text = text; r.append(text_node)
    h.append(r); paragraph._p.append(h)

def rich(paragraph, text):
    pattern = re.compile(r'(\[[^\]]+\]\((?:<[^>]+>|[^)]+)\)|\*\*[^*]+\*\*|`[^`]+`)')
    for token in pattern.split(text):
        if not token: continue
        if token.startswith('['):
            m = re.fullmatch(r'\[([^\]]+)\]\((.*)\)', token)
            hyperlink(paragraph, m[1], m[2])
        elif token.startswith('**'):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith('`'):
            run = paragraph.add_run(token[1:-1]); run.font.name = 'Consolas'; run.font.size = Pt(10)
        else:
            paragraph.add_run(token)

def table(lines):
    rows = [[x.strip() for x in line.strip().strip('|').split('|')] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r':?-+:?', cell) for cell in row)]
    n = len(rows[0]); t = doc.add_table(rows=0, cols=n); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    proportions = {3: [.24, .35, .41], 4: [.12, .3, .29, .29], 5: [.13, .215, .215, .22, .22]}.get(n, [1/n]*n)
    for col, width in zip(t.columns, proportions): col.width = Inches(6.9 * width)
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        b = OxmlElement('w:' + edge); b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '4'); b.set(qn('w:color'), 'D9D9D9'); borders.append(b)
    t._tbl.tblPr.append(borders)
    for idx, values in enumerate(rows):
        row = t.add_row()
        for cell, value, width in zip(row.cells, values, proportions):
            cell.width = Inches(6.9 * width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge in ['top', 'left', 'bottom', 'right']:
                m = OxmlElement('w:' + edge); m.set(qn('w:w'), '95'); m.set(qn('w:type'), 'dxa'); margins.append(m)
            pr.append(margins)
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E5EEF4' if idx == 0 else ('F6F7F8' if idx % 2 == 0 else 'FFFFFF')); pr.append(shade)
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(1); p.paragraph_format.line_spacing = 1.03
            if idx == 0: p.paragraph_format.keep_with_next = True
            rich(p, value)
            for run in p.runs:
                run.font.size = Pt(10.3)
                if idx == 0: run.bold = True
        if idx == 0:
            repeat = OxmlElement('w:tblHeader'); row._tr.get_or_add_trPr().append(repeat)
        no_split = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

header = section.header.paragraphs[0]
header.add_run('Nos Deniers  •  Audit et plan de développement').font.size = Pt(9)
footer = section.footer.paragraphs[0]
footer.add_run('9 septembre 2026  •  ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
for run in footer.runs: run.font.size = Pt(9)

lines = (OUT / 'AUDIT-NOS-DENIERS.md').read_text(encoding='utf-8').splitlines()
i = 0
while i < len(lines):
    line = lines[i].strip()
    if not line: i += 1; continue
    if line.startswith('|'):
        collected = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            collected.append(lines[i]); i += 1
        table(collected); continue
    if line.startswith('# '):
        doc.add_paragraph(line[2:], 'Title')
    elif line.startswith('## '):
        doc.add_heading(line[3:], 1)
    elif line.startswith('### '):
        doc.add_heading(line[4:], 2)
    elif re.match(r'^\d+\. ', line):
        p = doc.add_paragraph(style='List Number'); rich(p, re.sub(r'^\d+\. ', '', line))
    elif line.startswith('- '):
        p = doc.add_paragraph(style='List Bullet'); rich(p, line[2:])
    else:
        p = doc.add_paragraph(); rich(p, line)
    i += 1
for root in [doc.styles.element, doc.element]:
    for b in list(root.iter(qn('w:pBdr'))): b.getparent().remove(b)
doc.core_properties.title = 'Audit de Nos Deniers et plan de développement'
doc.core_properties.author = 'LexMachine'
doc.core_properties.subject = 'État vérifié du corpus et feuille de route du suivi budgétaire'
output = OUT / 'Audit Nos Deniers et plan de développement.docx'
doc.save(output)
print(output)
