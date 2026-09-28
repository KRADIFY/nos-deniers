"""Final audit memo, based on persisted checks, with explicit residual limits."""
from pathlib import Path
import json, urllib.request
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'reports/audit-final-20260920'
TARGET=OUT/'Nos Deniers - Audit final et limites de la livraison.docx'
REF=Path('C:/Users/Jean-Christophe/Desktop/Nos Deniers/Audit Nos Deniers et plan de développement.docx')
def read(p):return json.loads(Path(p).read_text(encoding='utf-8-sig'))
audit=read(OUT/'numeric-audit-final.json');files=read(OUT/'public-files-audit.json')
coverage=read(ROOT/'budget_service/data/rap-coverage.json')
assert audit['success'] and files['success']
d=Document(REF)
body=d._element.body
for child in list(body):
    if child.tag!=qn('w:sectPr'):body.remove(child)
sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=Cm(1.8);sec.left_margin=sec.right_margin=Cm(2)
normal=d.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.08
for name,size in [('Title',25),('Heading 1',17),('Heading 2',12)]:
    s=d.styles[name];s.font.color.rgb=RGBColor(0,0,0);s.font.size=Pt(size);s.font.bold=True
    s.paragraph_format.space_before=Pt(12);s.paragraph_format.space_after=Pt(8)
    for color in s.element.xpath('.//w:color'):
        for attr in ('themeColor','themeTint','themeShade'):color.attrib.pop(qn('w:'+attr),None)
for s in d.sections:
    s.header.paragraphs[0].text='Nos Deniers  •  Audit de livraison'
    s.footer.paragraphs[0].text='21 septembre 2026  •  '
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');s.footer.paragraphs[0]._p.append(field)

def p(text='',style=None):return d.add_paragraph(text,style)
def h(text):d.add_heading(text,level=1)
def sub(text):d.add_heading(text,level=2)
def page(title):d.add_page_break();h(title)
def link(par,label,url):
    rel=par.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    el=OxmlElement('w:hyperlink');el.set(qn('r:id'),rel);r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
    co=OxmlElement('w:color');co.set(qn('w:val'),'173B53');pr.append(co)
    u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u);r.append(pr)
    text=OxmlElement('w:t');text.text=label;r.append(text);el.append(r);par._p.append(el)
def table(headers,rows,widths):
    t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,width in zip(t.columns,widths):col.width=Cm(width)
    for cell,label in zip(t.rows[0].cells,headers):cell.text=str(label)
    trpr=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for cell,value in zip(cells,row):cell.text=str(value)
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
    t._tbl.tblPr.append(borders)
    for ri,row in enumerate(t.rows):
        for ci,cell in enumerate(row.cells):
            cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'173B53' if ri==0 else ('F1F5F8' if ri%2==0 else 'FFFFFF'));props.append(shade)
            margins=OxmlElement('w:tcMar')
            for side,n in [('top',85),('bottom',85),('left',100),('right',100)]:
                x=OxmlElement('w:'+side);x.set(qn('w:w'),str(n));x.set(qn('w:type'),'dxa');margins.append(x)
            props.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.03
                if ci==0:para.alignment=WD_ALIGN_PARAGRAPH.LEFT
                for run in para.runs:
                    run.font.size=Pt(10);run.font.bold=ri==0;run.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
    p()
    return t

d.add_heading('Nos Deniers\nAudit final de la livraison',0)
p('21 septembre 2026 · Pour Jean-Christophe et l’utilisatrice parlementaire')
h('Une version exploitable, avec des limites visibles')
p('La version locale a été enrichie et contrôlée. Elle permet de comparer les crédits proposés, votés et consommés, de retirer les postes identifiés, de corriger l’inflation et d’ouvrir les justificatifs. Les deux parcours prioritaires restent Mission Écologie et MaPrimeRénov’.')
p('Le paquet de publication est préparé et vérifié. Il n’a pas été activé sur le serveur : le site public sert encore la version précédente. Le présent document décrit la nouvelle version locale et la livraison du 21 septembre.')
table(['Résultat','État vérifié'],[
 ('Calculs et filtres','218 tests Python réussis ; parcours navigateur, exports et justificatifs contrôlés.'),
 ('Base de calcul','120 994 faits canoniques, complétés par des registres RAP et 68 observations MaPrimeRénov’.'),
 ('Corpus documentaire','4 439 documents / 4 051 433 passages dans la recherche hybride ; aucun recalcul des vecteurs.'),
 ('Anciennes extractions écartées','Les lots de 128 groupes d’actions, 223 pages de mouvements et 16 tables de réserves ont été retraités.'),
 ('Limites persistantes','Ventilations MaPrimeRénov’, certains tableaux RAP, historique de gestion, nomenclatures et fonctions avancées : détail dans ce mémo.')],[4.5,12.5])
p('Le verdict ne correspond pas à une garantie de couverture à 100 %. Des tableaux existent dans les archives sans être intégrés aux calculs. D’autres montants ne sont pas séparables dans les sources examinées. Ces deux situations sont distinguées dans les avertissements et ci-dessous.')
link(p(),'Ouvrir la version locale','http://127.0.0.1:8552/')

page('1. Ce qui a été contrôlé et corrigé')
p('Le contrôle a porté sur la base réellement servie par Docker, les registres complémentaires, leurs documents justificatifs et les parcours utilisés dans le navigateur. La lecture de passages vectorisés a servi à retrouver des sources ; elle n’a pas été assimilée à une validation des chiffres.')
table(['Contrôle','Résultat'],[
 ('Intégrité SQLite','Contrôle d’intégrité réussi ; conservation des 120 994 faits canoniques.'),
 ('Sources numériques','209 fichiers sources contrôlés par empreinte ; 25 216 références de pages vérifiées.'),
 ('Fichiers de livraison','4 291 fichiers contrôlés par taille et empreinte ; 3 892 références PDF avec en-tête valide. Ce dernier test n’est pas une relecture de chaque page.'),
 ('Rapprochements','264 rapprochements de missions 2024–2025 ; contrôles des actions, mouvements, réserves et crédits 2026.'),
 ('Comportements','AE/CP, inflation, exclusions et réactivation, cellules indisponibles, justificatifs, CSV/XLSX, recherche et affichage mobile.'),
 ('Protection technique','Lecture seule des données, utilisateur Docker sans privilèges, paramètres et chemins invalides refusés, en-têtes de sécurité vérifiés.')],[4.5,12.5])
sub('Corrections de cette dernière passe')
p('Douze observations historiques MaPrimeRénov’ ont été ajoutées, sans changer les 56 précédentes. Le retrait du dispositif ne masque plus les actes des programmes qui ne le financent pas. Les actes non ventilables ne fournissent plus un montant nominal susceptible d’être pris pour celui de la sélection. Les avertissements de couverture 2026 et l’aide aux interrupteurs ont été actualisés.')
p('Les écarts de l’identité « LFI + mouvements = ouverts » sont désormais expliqués dans le justificatif des lignes concernées. Les chiffres officiels restent conservés : le site n’invente pas un montant corrigé.')
p('La charte graphique, les logos et l’organisation générale des pages ont été conservés. Les tests de sécurité effectués sont ciblés ; ce bilan ne vaut pas certification de sécurité ni test d’intrusion exhaustif.')

page('2. La couverture annuelle réelle')
p('Le tableau donne le nombre de programmes distincts représentés en CP dans les tables du budget général. Il mesure la présence de lignes ; il ne prouve pas, à lui seul, une couverture exhaustive ni une nomenclature constante. Les changements de programmes expliquent une partie des différences entre colonnes.')
rows=[]
for year in range(2017,2027):
    def count(stage):return next((r['programmes'] for r in audit['coverage'] if (r['year'],r['budget'],r['measure'],r['stage'])==(year,'BG','CP',stage)), '—')
    rows.append((year,count('PLF'),count('LFI'),count('EXEC'),'Étapes détaillées 2023–2025' if 2023<=year<=2025 else 'Exercice en cours' if year==2026 else 'Gestion détaillée non généralisée'))
table(['Année','PLF','LFI','Consommé','Lecture'],rows,[1.6,1.5,1.5,2.1,10.3])
p('2026 : la LFI comprend aussi 5 programmes de budgets annexes, 15 de comptes d’affectation spéciale et 31 de comptes de concours financiers. Les 64 programmes BG comportant des FdC/AdP prévus au PLF ne suffisent pas à certifier ce détail pour tout le budget.')
p('Les budgets annexes, CAS et CCF se consultent séparément. La campagne récente de mouvements et réserves RAP porte sur le budget général 2023–2025 ; elle ne doit pas être présentée comme une couverture de tous les comptes spéciaux ni de toutes les années.')
p('Les cellules générales de gestion 2017–2022 restent une limite d’intégration. Des annexes publiques existent, notamment la situation par programme et titre des crédits ouverts et des dépenses 2021. Leur existence interdit de qualifier tout cet historique d’introuvable.')
link(p(),'Exemple officiel : annexes au règlement 2021','https://www.budget.gouv.fr/documentation/file-download/15361')

page('3. MaPrimeRénov’ : disponibilité et périmètres')
p('Les états ci-dessous portent sur le total du dispositif dans le périmètre publié par les sources, en crédits de l’État. Ils n’additionnent pas les versements de l’État à l’Anah et les paiements de l’Anah aux bénéficiaires. « Partiel » signifie que certaines parts sont visibles dans les justificatifs, mais ne remplacent pas le total.')
table(['Année','PLF','LFI','Ouverts','Consommé'],[
 ('2020','AE et CP','AE et CP','AE et CP','AE et CP'),
 ('2021','AE et CP','AE et CP','Partiel','AE et CP'),
 ('2022','AE et CP','Partiel','Partiel','AE et CP'),
 ('2023','Partiel','Partiel','Non établi','AE et CP'),
 ('2024','Partiel','AE et CP','AE et CP','AE et CP'),
 ('2025','CP seul','Partiel','Enveloppe mixte','Enveloppe mixte'),
 ('2026','Non établi','Partiel','Non établi','Exercice en cours')],[1.6,3.6,3.6,4.0,4.2])
sub('Ce que la dernière recherche a permis de compléter')
p('En 2020, la prime disposait de 390 M€ initiaux, puis de 575 M€ après 100 M€ supplémentaires et le transfert de 85 M€ depuis le P135. Ces 85 M€ sont déjà compris au P174 : ils ne sont pas comptés une seconde fois. Le RAP publie 575 M€ consommés en AE et 455 M€ en CP.')
p('Pour 2021, les PLF et LFI totalisent 2 740 M€ en AE et 1 655 M€ en CP. Pour le PLF 2022, le tableau parlementaire établit 1 700 M€ en AE et 1 955,5 M€ en CP. La colonne PLF n’est jamais recopiée dans la LFI et le collectif 2021 n’est pas traité comme le total annuel des ouverts.')
link(p(),'Assemblée nationale : tableau 2020–2022, pages 31–32','https://www.assemblee-nationale.fr/dyn/15/rapports/cion_fin/l15b4524-tiii-a45_rapport-fond.pdf#page=32')
p('Attention au périmètre : le PAP 2021 de relance inclut des aides connexes à la rénovation, les copropriétés et de la communication. La série suit cette définition publiée et ne constitue donc pas une série homogène de la seule prime par geste. Le CITE et les autres financements de rénovation restent distincts.')

page('4. Ce qui manque encore pour MaPrimeRénov’')
sub('Des enveloppes mixtes, pas un simple PDF manquant')
p('Les étapes nationales restantes ne peuvent pas être complétées en reprenant le total d’une action ou le budget entier de l’Anah. Les rapports examinés regroupent, selon l’année, plusieurs aides, les frais de fonctionnement, les aides aux copropriétés et d’autres interventions. Les définitions évoluent également entre prime par geste et rénovation globale.')
p('Pour 2025, le PLF publie 1 378 M€ de CP dans le périmètre de la prime de transition énergétique. Une brique de LFI de 779,9 M€ CP est identifiée ; elle ne représente pas le total, car une subvention générale à l’Anah finance aussi MaPrimeRénov’ dans des proportions non définies. La note 116 de la Cour des comptes explique cette limite. Pour 2026, la brique identifiée en LFI est de 604,6 M€ CP, avec la même difficulté de complément.')
link(p(),'Cour des comptes : Cohésion des territoires 2025, p. 55, note 116','https://www.ccomptes.fr/sites/default/files/2026-04/NEB-2026-Cohesion-territoires.pdf#page=55')
p('Le tableau 11 de cette note et le RAP sont accessibles, mais les enveloppes Anah qu’ils décrivent ne suffisent pas à isoler toutes les étapes du seul dispositif. Ces montants sont présentés comme contexte dans le justificatif, sans devenir automatiquement des valeurs à soustraire.')
sub('Pistes réexaminées')
p('Les recherches ont combiné le corpus local et la base Cour des comptes, les PAP et RAP des programmes 174, 362 et 135, les jaunes sur la rénovation énergétique, les notes d’exécution, ainsi que les rapports de l’Assemblée nationale et du Sénat. Les catalogues API de Bercy, data.gouv.fr et les publications Anah ont aussi été consultés pendant la campagne. Les nouveaux tableaux parlementaires ont été téléchargés et confrontés aux documents budgétaires.')
p('Le rapport parlementaire PLF 2023 comporte par exemple une ligne de 565 M€ CP en LFI 2022 au P362, contre 565,5 M€ au PLF 2022 ; son tableau agrège aussi la subvention Anah et présente des totaux qui demandent rapprochement. Il n’a pas été utilisé pour remplir silencieusement une LFI nationale complète.')
sub('Comportement dans le site')
p('Cliquer sur un montant montre son périmètre et ses sources. Cliquer sur une case indisponible explique ce qui est connu, ce qui est mélangé et pourquoi le retrait serait trompeur. Les gels et mouvements ne sont pas ventilés automatiquement au prorata de MaPrimeRénov’. Aucune demande à une administration n’a été envoyée dans cet audit.')

page('5. Actions, mouvements et réserves RAP')
table(['Bloc','Ancien état cité','État de la livraison'],[
 ('Actions RAP','1 231 groupes nationaux ; 128 écartés','1 401 groupes nationaux + 115 pilotes conservés : 1 516 groupes, 7 900 lignes d’actions et 3 010 de sous-actions.'),
 ('Mouvements RAP','126 programmes-années ; 223 pages candidates non promues','354 programmes-années nationaux + 3 pilotes P174, soit 357 sur 379 audités en BG 2023–2025.'),
 ('Réserves RAP','264 programmes-années ; 16 tables ambiguës','303 programmes-années nationaux + 28 Écologie, soit 331 sur 379 audités en BG 2023–2025. En plus : 48 Écologie de 2017 à 2022.')],[3.2,5.2,8.6])
p('Les nombres de lignes d’actions sont comptés par année, étape et mesure AE/CP ; il ne s’agit pas d’autant de politiques distinctes. Aucun rejet d’extraction ne subsiste dans les trois lots de candidats cités. Cela ne signifie pas que tous les programmes possèdent les tableaux recherchés.')
p('Sur les 379 programmes-années audités en 2023–2025, 22 n’ont pas de récapitulation de mouvements intégrée et 48 n’ont pas de tableau de réserve. Ces listes se recoupent et représentent 53 programmes-années distincts, détaillés en annexe. Pour le P384 en 2025, l’absence de réserve est expressément documentée : le site affiche « Sans objet » plutôt qu’un zéro inventé.')
sub('Les écarts sont documentés sans faire disparaître les chiffres')
p('Les 97 notes d’identité de mouvements sont conservées. Le plus grand écart de cette série est de 402 462 € pour le P200 en 2023, sur 139 544 000 000 € ouverts en AE, soit environ 0,00029 %. Les valeurs source sont conservées. Autre cas connu : P124 en 2023, 1 226 538 243,70 € dans la synthèse contre 1 227 063 244 € au RAP p. 137 ; l’écart de 525 000,30 € est exposé, sans double compte.')
p('Ces exemples ne constituent pas une déclaration du plus grand écart de tous les documents existants. Les rapprochements distinguent arrondis, divergences de source et changements de périmètre ; un petit écart n’est plus un motif automatique de rejet.')
sub('Ce qu’un RAP annuel ne permet pas de déduire')
p('Mise en réserve, surgels, dégels, annulations de réserve et solde sont distingués. Le non-consommé n’est pas assimilé au gel. Les tableaux annuels ne donnent pas systématiquement une chronologie complète de chaque décision, ni une ventilation par action. Le registre juridique d’actes demeure partiel ; les tableaux RAP ne sont pas additionnés une seconde fois aux crédits ouverts.')

page('6. Recherche, application et mise en ligne')
p('La recherche hybride locale fonctionne sur 4 439 documents et 4 051 433 passages. Le navigateur a été testé avec une recherche réelle, le filtre d’année, une page de source, l’ouverture du passage complet et le contrôle de son empreinte. Les sources ajoutées après la vectorisation sont accessibles depuis les justificatifs et la bibliothèque ; elles ne sont pas automatiquement présentes dans cet index figé.')
p('La connexion LexMachine à la base Cour des comptes a fourni des résultats utiles. Les appels Tricoteuses testés ont renvoyé une erreur 401 évoquant la clé interne x-typesense-api-key. Ce résultat signale un connecteur à réparer ; il ne démontre ni un refus de votre OAuth ni une indisponibilité des publications officielles. Les recherches locales et les accès web officiels ont permis de poursuivre.')
sub('Fonctions effectivement livrées')
p('Tableaux et graphiques, AE/CP, comparaisons PLF/LFI/consommé, gestion récente, exclusions avec interrupteurs individuels, correction annuelle de l’inflation avant cumul, formats euro/million/milliard, séparateurs de milliers, exports CSV/XLSX/JSON et justificatifs cliquables. Les sélections sont mémorisées dans le navigateur.')
sub('Fonctions qui ne doivent pas être annoncées comme achevées')
p('Une analyse causale automatique par IA n’est pas livrée : la recherche fournit les pièces, mais n’explique pas à elle seule pourquoi une politique a réussi ou échoué. La navette parlementaire, les espaces collaboratifs et les sélections partagées ne font pas partie de cette livraison. La correspondance intégrale des nomenclatures 2017–2026 et la gestion détaillée historique hors Écologie restent à construire ou à intégrer.')
sub('Livraison Docker')
p('La livraison 20260921-audit-final a été vérifiée hors réseau puis par HTTP en local. Elle comprend l’image web, les différences de données, les nouveaux PDF, les inventaires et un installateur contrôlé. Le retour à la version publique précédente est prévu si la vérification après activation échoue. Seul le service Budget est ciblé.')
p('Le premier transfert complet de la recherche représente environ 33,3 Go. Les index et l’image de recherche déjà préparés sont réutilisés, sans revectorisation ni duplication locale de tout le corpus. Le transfert SFTP reprend les fichiers partiels et vérifie les empreintes. Les identifiants personnels des connecteurs ne sont pas inclus.')

page('7. Conclusion de l’audit et limites des recherches')
p('Le site est livrable pour travailler sur les données intégrées et leurs justificatifs. Mission Écologie permet des comparaisons fines sur les postes documentés ; MaPrimeRénov’ permet d’isoler ou retirer les seules parts établies, en conservant les avertissements de définition. Les contrôles réussis ne rendent pas comparables des enveloppes de périmètres différents.')
table(['Nature du manque','Ce que le bilan permet de conclure'],[
 ('Donnée non séparable','Certaines ventilations MaPrimeRénov’ 2022–2026 et leurs gels ne sont pas établies dans les sources examinées. Les enveloppes mixtes restent du contexte.'),
 ('Tableau absent du RAP contrôlé','22 récapitulations de mouvements et 48 tableaux de réserves manquent sur le périmètre récent audité ; 1 réserve est explicitement sans objet.'),
 ('Donnée encore non intégrée','Des archives de gestion 2017–2022 existent. Elles demandent une extraction et un rapprochement spécifiques ; ce travail ne peut pas être déclaré épuisé.'),
 ('Temporalité','Le résultat annuel 2026 n’est pas encore clos. Une situation infra-annuelle doit porter sa date.'),
 ('Accès technique','Chorus n’est pas accessible ; Tricoteuses renvoie l’erreur constatée. Le site fonctionne avec le corpus local et les sources publiques.'),
 ('Fonction à développer','Audit narratif IA, nomenclatures homogènes complètes, navette et collaboration restent distincts de la qualité des données disponibles.')],[4.2,12.8])
p('La recherche a abouti à une borne de livraison documentée, pas à la preuve qu’aucune autre information ne peut exister. Les sources examinées et les limites d’intégration restent consignées pour permettre une reprise sans recommencer les contrôles déjà réussis.')
sub('Accès aux résultats')
link(p(),'Version locale auditée','http://127.0.0.1:8552/')
link(p(),'Site public — ancienne version tant que la livraison n’est pas activée','https://budget.lexmachine.net/')
p('Dossier de livraison : deploy/update-20260921-audit-final dans le projet Budget. Le fichier LIRE_AVANT_LIVRAISON.txt explique le transfert et la commande serveur. Les preuves techniques sont conservées dans reports/audit-final-20260920 ; le point de reprise est POINT_DE_REPRISE-AUDIT-FINAL-20260921.md.')

gaps=coverage['gaps']
for part, (offset, end) in enumerate(((0, 18), (18, 35), (35, len(gaps))), 1):
    page(f'Annexe — tableaux RAP restant à documenter ({part}/3)')
    p('Périmètre : budget général 2023–2025. « — » indique le tableau non intégré après contrôle du RAP ; « présent » indique que le programme n’appartient pas à la liste des lacunes de ce bloc. La liste n’inclut pas les années antérieures ni les autres budgets.')
    chunk=gaps[offset:end]
    rows=[]
    for r in chunk:
        reserve='Sans objet' if (r['year'],r['program'])==(2025,'384') else ('—' if 'reserves' in r['missing'] else 'Présent')
        rows.append((r['year'],r['program']+' · '+r['program_label'],reserve,'—' if 'movements' in r['missing'] else 'Présent',str(r['page'])))
    t=table(['Année','Programme','Réserve','Mouvements','Page PDF'],rows,[1.4,9.1,2.1,2.4,2.0])
    for row in t.rows:
        for cell in row.cells:
            for margin in cell._tc.xpath('.//w:tcMar/w:top | .//w:tcMar/w:bottom'):
                margin.set(qn('w:w'), '55')
    for i,r in enumerate(chunk,1):
        par=t.rows[i].cells[-1].paragraphs[0];par.clear();link(par,str(r['page']),'http://127.0.0.1:8552/api/download/'+r['source']+'#page='+str(r['page']))
    p('Les liens de page ouvrent les copies de la version locale. Les justificatifs du site précisent la portée du manque, les références et les éventuelles exemptions documentées.')

d.core_properties.title='Nos Deniers — Audit final et limites de la livraison'
d.core_properties.subject='Livraison du 21 septembre 2026, couverture et limites vérifiées'
d.core_properties.author='Nos Deniers · LexMachine'
d.save(TARGET)
print(TARGET)
